# -*- coding: utf-8 -*-
"""
Add a final class column to processed section workbooks.

Class labels are assigned per run by matching each segment runtime to the
nearest Class1-Class5 standard runtime for that section.
"""
from __future__ import annotations

import argparse
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.writer.excel import ExcelWriter


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT_DIR = PROJECT_ROOT / "data" / "data_processed"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data" / "data_processed_with_class"
DEFAULT_STANDARD_TIMES = (
    PROJECT_ROOT / "output" / "analysis" / "class_tables_strict" / "standard_class_times.csv"
)
FAST_XLSX_COMPRESSLEVEL = 1

TIME_COL = "时刻"
SEGMENT_COL = "segment"
SECTION_COL = "区段"
CLASS_TIME_COLS = ["Class1", "Class2", "Class3", "Class4", "Class5"]


@dataclass(frozen=True)
class ClassMatch:
    section_pair: str
    matched_pair: str
    segment: object
    runtime_s: float
    class_label: str
    standard_time_s: float
    abs_gap_s: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Append a final runtime-nearest class column to processed results_*.xlsx files."
    )
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--standard-times", type=Path, default=DEFAULT_STANDARD_TIMES)
    parser.add_argument("--pattern", default="results_*.xlsx")
    parser.add_argument("--class-col", default="class", help="Output class column name.")
    parser.add_argument(
        "--in-place",
        action="store_true",
        help="Overwrite input files instead of writing to --output-dir.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Only process the first N matched files, useful for checking.",
    )
    return parser.parse_args()


def reverse_pair(pair: str) -> str:
    parts = str(pair).split("-")
    if len(parts) != 2:
        return str(pair)
    return f"{parts[1]}-{parts[0]}"


def infer_section_pair(path: Path, df: pd.DataFrame) -> str:
    if SECTION_COL in df.columns:
        vals = df[SECTION_COL].dropna().astype(str).str.strip()
        if not vals.empty:
            return vals.iloc[0]

    stem = path.stem
    for prefix in ("results_", "cleaned_"):
        if stem.startswith(prefix):
            return stem[len(prefix) :]
    return stem


def load_standard_times(path: Path) -> dict[str, dict[str, float]]:
    if not path.exists():
        raise FileNotFoundError(f"standard class time file not found: {path}")

    table = pd.read_csv(path)
    missing = [c for c in [SECTION_COL, *CLASS_TIME_COLS] if c not in table.columns]
    if missing:
        raise ValueError(f"standard class time file missing columns: {missing}")

    class_times: dict[str, dict[str, float]] = {}
    for _, row in table.iterrows():
        pair = str(row[SECTION_COL]).strip()
        class_times[pair] = {
            c.lower(): float(row[c])
            for c in CLASS_TIME_COLS
            if pd.notna(row[c])
        }
    return class_times


def get_section_class_times(
    section_pair: str, class_times: dict[str, dict[str, float]]
) -> tuple[str, dict[str, float]]:
    if section_pair in class_times:
        return section_pair, class_times[section_pair]

    reversed_pair = reverse_pair(section_pair)
    if reversed_pair in class_times:
        return reversed_pair, class_times[reversed_pair]

    raise KeyError(f"no class standard time found for {section_pair} or {reversed_pair}")


def build_runtime_matches(
    df: pd.DataFrame,
    section_pair: str,
    matched_pair: str,
    standards: dict[str, float],
) -> tuple[pd.Series, list[ClassMatch]]:
    if TIME_COL not in df.columns:
        raise ValueError(f"missing required time column: {TIME_COL}")

    work = df.copy()
    if SEGMENT_COL not in work.columns:
        work[SEGMENT_COL] = build_segment_from_time_reset(work[TIME_COL])

    time_values = pd.to_numeric(work[TIME_COL], errors="coerce")
    segment_values = work[SEGMENT_COL]
    runtime_by_segment = time_values.groupby(segment_values).agg(lambda s: s.max() - s.min())

    labels: dict[object, str] = {}
    matches: list[ClassMatch] = []
    class_labels = list(standards.keys())
    class_seconds = np.array([standards[c] for c in class_labels], dtype=float)

    for segment, runtime in runtime_by_segment.items():
        if pd.isna(runtime):
            continue
        gaps = np.abs(class_seconds - float(runtime))
        best_idx = int(np.argmin(gaps))
        best_class = class_labels[best_idx]
        best_standard = float(class_seconds[best_idx])
        labels[segment] = best_class
        matches.append(
            ClassMatch(
                section_pair=section_pair,
                matched_pair=matched_pair,
                segment=segment,
                runtime_s=round(float(runtime), 3),
                class_label=best_class,
                standard_time_s=best_standard,
                abs_gap_s=round(float(gaps[best_idx]), 3),
            )
        )

    class_series = segment_values.map(labels)
    return class_series, matches


def build_runtime_matches_from_rows(
    times: list[object],
    segments: list[object],
    section_pair: str,
    matched_pair: str,
    standards: dict[str, float],
) -> tuple[list[str | None], list[ClassMatch]]:
    if segments:
        segment_values = pd.Series(segments)
    else:
        segment_values = build_segment_from_time_reset(pd.Series(times))

    time_values = pd.to_numeric(pd.Series(times), errors="coerce")
    runtime_by_segment = time_values.groupby(segment_values).agg(lambda s: s.max() - s.min())

    labels: dict[object, str] = {}
    matches: list[ClassMatch] = []
    class_labels = list(standards.keys())
    class_seconds = np.array([standards[c] for c in class_labels], dtype=float)

    for segment, runtime in runtime_by_segment.items():
        if pd.isna(runtime):
            continue
        gaps = np.abs(class_seconds - float(runtime))
        best_idx = int(np.argmin(gaps))
        best_class = class_labels[best_idx]
        best_standard = float(class_seconds[best_idx])
        labels[segment] = best_class
        matches.append(
            ClassMatch(
                section_pair=section_pair,
                matched_pair=matched_pair,
                segment=segment,
                runtime_s=round(float(runtime), 3),
                class_label=best_class,
                standard_time_s=best_standard,
                abs_gap_s=round(float(gaps[best_idx]), 3),
            )
        )

    return [labels.get(segment) for segment in segment_values], matches


def build_segment_from_time_reset(time_col: pd.Series) -> pd.Series:
    t = pd.to_numeric(time_col, errors="coerce").ffill().fillna(0.0).to_numpy(dtype=float)
    reset = np.diff(t, prepend=t[0]) < 0
    return pd.Series(np.cumsum(reset), index=time_col.index)


def reorder_with_final_class(df: pd.DataFrame, class_col: str, values: pd.Series) -> pd.DataFrame:
    out = df.copy()
    if class_col in out.columns:
        out = out.drop(columns=[class_col])
    out[class_col] = values.values
    return out


def header_index(headers: list[object], name: str) -> int | None:
    try:
        return headers.index(name)
    except ValueError:
        return None


def first_non_empty(values: Iterable[object]) -> str | None:
    for value in values:
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def analyze_workbook(
    source: Path,
    class_times: dict[str, dict[str, float]],
) -> tuple[dict[str, list[str | None]], list[ClassMatch]]:
    wb = load_workbook(source, read_only=True, data_only=True)
    class_values_by_sheet: dict[str, list[str | None]] = {}
    all_matches: list[ClassMatch] = []

    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows, []) or [])
            time_idx = header_index(headers, TIME_COL)
            if time_idx is None:
                continue

            segment_idx = header_index(headers, SEGMENT_COL)
            section_idx = header_index(headers, SECTION_COL)
            times: list[object] = []
            segments: list[object] = []
            section_values: list[object] = []

            for row in rows:
                if time_idx >= len(row):
                    times.append(None)
                else:
                    times.append(row[time_idx])

                if segment_idx is not None:
                    segments.append(row[segment_idx] if segment_idx < len(row) else None)

                if section_idx is not None:
                    section_values.append(row[section_idx] if section_idx < len(row) else None)

            section_pair = first_non_empty(section_values) or infer_section_pair(source, pd.DataFrame())
            matched_pair, standards = get_section_class_times(section_pair, class_times)
            class_values, matches = build_runtime_matches_from_rows(
                times,
                segments,
                section_pair,
                matched_pair,
                standards,
            )
            class_values_by_sheet[ws.title] = class_values
            all_matches.extend(matches)
    finally:
        wb.close()

    return class_values_by_sheet, all_matches


def stream_write_with_class(
    source: Path,
    destination: Path,
    class_values_by_sheet: dict[str, list[str | None]],
    class_col: str,
) -> None:
    in_wb = load_workbook(source, read_only=True, data_only=False)
    out_wb = Workbook(write_only=True)
    if out_wb.worksheets:
        out_wb.remove(out_wb.worksheets[0])

    try:
        for in_ws in in_wb.worksheets:
            out_ws = out_wb.create_sheet(title=in_ws.title)
            values_iter = in_ws.iter_rows(values_only=True)
            headers = list(next(values_iter, []) or [])
            existing_class_idx = header_index(headers, class_col)
            has_class_values = in_ws.title in class_values_by_sheet

            out_headers = [
                value
                for idx, value in enumerate(headers)
                if existing_class_idx is None or idx != existing_class_idx
            ]
            if has_class_values:
                out_headers.append(class_col)
            out_ws.append(out_headers)

            class_values = class_values_by_sheet.get(in_ws.title, [])
            for row_zero, row in enumerate(values_iter):
                out_row = [
                    value
                    for idx, value in enumerate(row)
                    if existing_class_idx is None or idx != existing_class_idx
                ]
                if has_class_values:
                    value = class_values[row_zero] if row_zero < len(class_values) else None
                    out_row.append(value)
                out_ws.append(out_row)

        destination.parent.mkdir(parents=True, exist_ok=True)
        save_workbook_fast(out_wb, destination)
    finally:
        in_wb.close()
        out_wb.close()


def save_workbook_fast(workbook: Workbook, destination: Path) -> None:
    with ZipFile(
        destination,
        "w",
        ZIP_DEFLATED,
        allowZip64=True,
        compresslevel=FAST_XLSX_COMPRESSLEVEL,
    ) as archive:
        ExcelWriter(workbook, archive).save()


def write_workbook(
    source: Path,
    destination: Path,
    class_times: dict[str, dict[str, float]],
    class_col: str,
) -> list[ClassMatch]:
    class_values_by_sheet, all_matches = analyze_workbook(source, class_times)
    if source.resolve() == destination.resolve():
        temp_destination = destination.with_suffix(".tmp.xlsx")
        stream_write_with_class(source, temp_destination, class_values_by_sheet, class_col)
        temp_destination.replace(destination)
    else:
        stream_write_with_class(source, destination, class_values_by_sheet, class_col)
    return all_matches


def iter_input_files(input_dir: Path, pattern: str, limit: int | None) -> Iterable[Path]:
    files = sorted(input_dir.glob(pattern))
    if limit is not None:
        files = files[:limit]
    return files


def main() -> None:
    args = parse_args()
    input_dir: Path = args.input_dir
    output_dir: Path = args.output_dir

    if not input_dir.exists():
        raise FileNotFoundError(f"input dir not found: {input_dir}")

    class_times = load_standard_times(args.standard_times)
    files = list(iter_input_files(input_dir, args.pattern, args.limit))
    if not files:
        raise FileNotFoundError(f"no files matched {args.pattern} in {input_dir}")

    all_matches: list[ClassMatch] = []
    failures: list[tuple[str, str]] = []
    for idx, source in enumerate(files, start=1):
        destination = source if args.in_place else output_dir / source.name
        print(f"[{idx}/{len(files)}] {source.name} -> {destination}")
        try:
            matches = write_workbook(source, destination, class_times, args.class_col)
            all_matches.extend(matches)
        except Exception as exc:
            failures.append((source.name, str(exc)))
            print(f"  failed: {exc}")

    summary_dir = input_dir if args.in_place else output_dir
    summary_dir.mkdir(parents=True, exist_ok=True)
    if all_matches:
        summary = pd.DataFrame([m.__dict__ for m in all_matches])
        summary.to_csv(summary_dir / "class_assignment_summary.csv", index=False, encoding="utf-8-sig")

    if failures:
        pd.DataFrame(failures, columns=["file", "error"]).to_csv(
            summary_dir / "class_assignment_failures.csv", index=False, encoding="utf-8-sig"
        )

    print(f"Done. files={len(files)}, segments={len(all_matches)}, failures={len(failures)}")
    print(f"Output dir: {summary_dir}")


if __name__ == "__main__":
    main()
