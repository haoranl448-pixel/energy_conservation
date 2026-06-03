# -*- coding: utf-8 -*-
"""统计每个区间在 Step2 结果中实际出现过哪些运行等级。

这个脚本面向 `results_*.xlsx` 这类 Step2 输出文件：
1. 每个 Excel 文件通常对应一个站间区间；
2. 一个文件里有多趟车，每趟车用 `segment` 区分；
3. 同一趟车有很多行轨迹点，统计等级时只按一趟算 1 次；
4. 如果文件里有 `曲线质量标签`，总表会额外统计“正常标签 0”保留下来的等级。
"""

from __future__ import annotations

import argparse
import math
import re
from datetime import datetime, time
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_first5_curve_quality"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output" / "analysis" / "section_class_distribution"

COL_SECTION = "区段"
COL_DIRECTION = "列车运行方向"
COL_SEGMENT = "segment"
COL_CLASS = "运行等级"
COL_QUALITY = "曲线质量标签"
COL_AFTER_8PM = "是否是晚上八点以后的趟次"
COL_TIME = "时刻"

CLASS_ORDER = ["class1", "class2", "class3", "class4", "class5"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="统计 Step2 results_*.xlsx 中每个区间实际出现过的运行等级。",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Step2 结果文件所在目录。")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="统计结果输出目录。")
    parser.add_argument("--pattern", default="results_*.xlsx", help="输入文件匹配模式。")
    parser.add_argument("--normal-label", type=int, default=0, help="曲线质量标签中代表正常白天曲线的数值。")
    return parser.parse_args()


def normalize_class_label(value: Any) -> str | None:
    """把 `3`、`3.0`、`Class3`、`等级3` 等写法统一成 `class3`。"""
    if pd.isna(value):
        return None

    text = str(value).strip().lower()
    if not text:
        return None

    match = re.search(r"(\d+(?:\.\d+)?)", text)
    if not match:
        return text

    num = float(match.group(1))
    if num.is_integer():
        return f"class{int(num)}"
    return text


def class_sort_key(label: str) -> tuple[int, str]:
    if label in CLASS_ORDER:
        return CLASS_ORDER.index(label), label
    return len(CLASS_ORDER), label


def class_list(series: pd.Series) -> str:
    labels = sorted({str(x) for x in series.dropna() if str(x)}, key=class_sort_key)
    return ",".join(labels)


def count_by_class(segment_df: pd.DataFrame, prefix: str) -> dict[str, int]:
    counts = segment_df[COL_CLASS].value_counts(dropna=False).to_dict()
    return {f"{prefix}_{cls}_趟数": int(counts.get(cls, 0)) for cls in CLASS_ORDER}


def compact_label_counts(series: pd.Series) -> str:
    if series.empty:
        return ""
    counts = series.value_counts(dropna=False).sort_index()
    return "; ".join(f"{int(idx)}:{int(val)}" for idx, val in counts.items())


def to_float(value: Any) -> float | None:
    """把 Excel 单元格值尽量转成浮点数，无法转换时返回 None。"""
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, datetime):
        return value.timestamp()
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    try:
        num = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(num):
        return None
    return num


def value_at(row: tuple[Any, ...], col_idx: int | None) -> Any:
    if col_idx is None or col_idx >= len(row):
        return None
    return row[col_idx]


def read_segment_summary(path: Path) -> pd.DataFrame:
    """读取单个区间文件，并把多行轨迹点压缩成一趟一行。

    这里不用 `pandas.read_excel`，因为当前文件比较大，pandas 会把整块
    数据拉进内存。`openpyxl` 的只读模式可以一行一行扫，只保留每个
    segment 的少量统计字段。
    """
    wanted_cols = {
        COL_SECTION,
        COL_DIRECTION,
        COL_SEGMENT,
        COL_CLASS,
        COL_QUALITY,
        COL_AFTER_8PM,
        COL_TIME,
    }
    fallback_section = path.stem.replace("results_", "", 1)
    rows_by_segment: dict[int, dict[str, Any]] = {}

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_map = {str(name).strip(): idx for idx, name in enumerate(header) if name is not None}
        col_idx = {col: header_map.get(col) for col in wanted_cols}

        missing = [col for col in [COL_SEGMENT, COL_CLASS] if col_idx.get(col) is None]
        if missing:
            raise ValueError(f"{path.name} 缺少必要列: {missing}")

        for row in sheet.iter_rows(min_row=2, values_only=True):
            seg_num = to_float(value_at(row, col_idx[COL_SEGMENT]))
            if seg_num is None:
                continue

            segment = int(seg_num)
            rec = rows_by_segment.setdefault(
                segment,
                {
                    COL_SEGMENT: segment,
                    "区段": fallback_section,
                    "列车运行方向": "",
                    "运行等级": None,
                    "曲线质量标签": pd.NA,
                    "是否是晚上八点以后的趟次": 0,
                    "起始时刻s": pd.NA,
                    "结束时刻s": pd.NA,
                    "来源文件": path.name,
                },
            )

            section = value_at(row, col_idx.get(COL_SECTION))
            if section not in (None, "") and rec["区段"] == fallback_section:
                rec["区段"] = str(section)

            direction = value_at(row, col_idx.get(COL_DIRECTION))
            if direction not in (None, "") and not rec["列车运行方向"]:
                rec["列车运行方向"] = str(direction)

            class_label = normalize_class_label(value_at(row, col_idx[COL_CLASS]))
            if class_label and rec["运行等级"] is None:
                rec["运行等级"] = class_label

            quality = to_float(value_at(row, col_idx.get(COL_QUALITY)))
            if quality is not None and pd.isna(rec["曲线质量标签"]):
                rec["曲线质量标签"] = int(quality)

            after_8pm = to_float(value_at(row, col_idx.get(COL_AFTER_8PM)))
            if after_8pm is not None:
                rec["是否是晚上八点以后的趟次"] = max(int(after_8pm), int(rec["是否是晚上八点以后的趟次"]))

            t_val = to_float(value_at(row, col_idx.get(COL_TIME)))
            if t_val is not None:
                if pd.isna(rec["起始时刻s"]) or t_val < rec["起始时刻s"]:
                    rec["起始时刻s"] = t_val
                if pd.isna(rec["结束时刻s"]) or t_val > rec["结束时刻s"]:
                    rec["结束时刻s"] = t_val
    finally:
        workbook.close()

    records = []
    for rec in sorted(rows_by_segment.values(), key=lambda item: item[COL_SEGMENT]):
        if pd.isna(rec["起始时刻s"]) or pd.isna(rec["结束时刻s"]):
            rec["运行时长s"] = pd.NA
        else:
            rec["运行时长s"] = rec["结束时刻s"] - rec["起始时刻s"]
        records.append(rec)

    return pd.DataFrame.from_records(records)


def build_section_summary(segment_df: pd.DataFrame, normal_label: int) -> pd.DataFrame:
    rows = []
    group_keys = ["区段", "列车运行方向", "来源文件"]

    for (section, direction, source_file), group in segment_df.groupby(group_keys, dropna=False):
        has_quality = group["曲线质量标签"].notna().any()
        normal = group[group["曲线质量标签"] == normal_label] if has_quality else group

        row = {
            "来源文件": source_file,
            "区段": section,
            "列车运行方向": direction,
            "总趟次": int(len(group)),
            "全部等级列表": class_list(group["运行等级"]),
            "正常标签": normal_label if has_quality else "",
            "正常趟次": int(len(normal)),
            "正常等级列表": class_list(normal["运行等级"]),
            "质量标签分布": compact_label_counts(group["曲线质量标签"].dropna().astype(int)),
            "20点以后趟次": int(group["是否是晚上八点以后的趟次"].fillna(0).astype(int).sum()),
        }
        row.update(count_by_class(group, "全部"))
        row.update(count_by_class(normal, "正常"))
        rows.append(row)

    out = pd.DataFrame(rows)
    if out.empty:
        return out
    return out.sort_values(["来源文件", "区段", "列车运行方向"]).reset_index(drop=True)


def write_outputs(summary: pd.DataFrame, detail: pd.DataFrame, output_dir: Path) -> tuple[Path, Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    summary_csv = output_dir / "section_class_distribution.csv"
    detail_csv = output_dir / "section_class_segment_detail.csv"
    workbook = output_dir / "section_class_distribution.xlsx"

    summary.to_csv(summary_csv, index=False, encoding="utf-8-sig")
    detail.to_csv(detail_csv, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="区间等级汇总", index=False)
        detail.to_excel(writer, sheet_name="趟次明细", index=False)

    return summary_csv, detail_csv, workbook


def main() -> int:
    args = parse_args()
    data_dir = Path(args.data_dir)
    output_dir = Path(args.output_dir)

    if not data_dir.exists() or not data_dir.is_dir():
        raise FileNotFoundError(f"数据目录不存在: {data_dir}")

    files = sorted(data_dir.glob(args.pattern))
    if not files:
        raise FileNotFoundError(f"未找到文件: {data_dir / args.pattern}")

    print(f"扫描数据目录: {data_dir}", flush=True)
    print(f"匹配文件数量: {len(files)}", flush=True)

    segment_tables = []
    for i, path in enumerate(files, start=1):
        print(f"[{i}/{len(files)}] {path.name}", flush=True)
        segment_tables.append(read_segment_summary(path))

    detail = pd.concat(segment_tables, ignore_index=True)
    summary = build_section_summary(detail, args.normal_label)
    summary_csv, detail_csv, workbook = write_outputs(summary, detail, output_dir)

    print("\n各区间等级分布:")
    preview_cols = ["区段", "列车运行方向", "总趟次", "全部等级列表", "正常趟次", "正常等级列表", "质量标签分布"]
    print(summary[preview_cols].to_string(index=False))
    print(f"\n汇总 CSV: {summary_csv}")
    print(f"明细 CSV: {detail_csv}")
    print(f"Excel 总表: {workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
