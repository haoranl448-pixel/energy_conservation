# -*- coding: utf-8 -*-
"""
Batch-build historical-only section reports for selected trips.

This script is intentionally lightweight compared with the main pipeline:
it reads the processed Step2 `results_<section>.xlsx` files and directly
summarizes measured historical running time, raw measured energy, speed,
distance, mass, class, and quality label. It does not run residual models,
energy menus, DP scheduling, or OpenTrack exports.
"""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_all_curve_quality"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output" / "schedule" / "batch_trip_reports_trip1_10"

FULL_LINE_STATIONS = [
    "布政-张家潭", "张家潭-同德路", "同德路-石碶", "石碶-雅渡", "雅渡-庙堰",
    "庙堰-钟公庙", "钟公庙-鄞州区政府", "鄞州区政府-钱湖南路", "钱湖南路-南高教园区",
    "南高教园区-下应路", "下应路-大洋江", "大洋江-泗港", "泗港-曹隘", "曹隘-柳隘",
    "柳隘-海晏北路", "海晏北路-民安东路", "民安东路-会展中心", "会展中心-院士路",
    "院士路-盎孟港", "盎孟港-三官堂", "三官堂-兴庄路", "兴庄路-兴海南路",
    "兴海南路-梅堰", "梅堰-永茂路", "永茂路-镇海大道", "镇海大道-骆驼桥",
]

SOURCE_COLS = [
    "区段",
    "站间距",
    "时段划分",
    "日期+服务号",
    "服务号",
    "车底号",
    "列车运行方向",
    "时刻",
    "segment",
    "energy",
    "速度(m/s)",
    "累计位移(m)",
    "重量",
    "运行等级",
    "是否是晚上八点以后的趟次",
    "曲线质量标签",
]

SECTION_COL = "站间区间"
RUN_TIME_COL = "历史运行时间(s)"
DWELL_COL = "历史停站时间(s)"
TOTAL_TIME_COL = "历史累计用时含停站(s)"
ENERGY_COL = "历史实测能耗(Wh)"
DISTANCE_COL = "区间距离(m)"
STATION_DISTANCE_COL = "站间距(m)"
AVG_SPEED_COL = "平均速度(km/h)"
MAX_SPEED_COL = "最高速度(km/h)"
SAMPLE_AVG_SPEED_COL = "速度点平均(km/h)"
SPEED_SAMPLE_COUNT_COL = "速度采样点数"
CRUISE_AVG_SPEED_COL = "巡航平均速度(km/h)"
CRUISE_POINT_COUNT_COL = "巡航点数"
P95_SPEED_COL = "P95速度(km/h)"
P99_SPEED_COL = "P99速度(km/h)"
MASS_COL = "平均重量(t)"


def parse_trip_list(text: str) -> list[int]:
    trips: list[int] = []
    for part in text.split(","):
        item = part.strip()
        if not item:
            continue
        if "-" in item:
            left_text, right_text = item.split("-", 1)
            left = int(left_text.strip())
            right = int(right_text.strip())
            if left <= 0 or right <= 0 or right < left:
                raise ValueError(f"Invalid trip range: {item}")
            trips.extend(range(left, right + 1))
        else:
            trip = int(item)
            if trip <= 0:
                raise ValueError("--trips must use 1-based positive trip numbers.")
            trips.append(trip)
    if not trips:
        raise ValueError("No trips selected.")
    return sorted(dict.fromkeys(trips))


def parse_csv_list(value: str | None) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def resolve_project_path(path_value: str | Path) -> Path:
    path = Path(path_value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def choose_sections(line_scope: str, sections_text: str | None) -> list[str]:
    sections = parse_csv_list(sections_text)
    if sections:
        return sections

    scope = str(line_scope).strip().lower()
    if scope in {"full", "all"}:
        return FULL_LINE_STATIONS
    if scope == "first5":
        return FULL_LINE_STATIONS[:5]

    try:
        n_sections = int(scope)
    except ValueError as exc:
        raise ValueError("--line-scope must be 'full' or a positive integer.") from exc
    if n_sections < 1:
        raise ValueError("--line-scope must be >= 1.")
    if n_sections > len(FULL_LINE_STATIONS):
        raise ValueError(f"--line-scope cannot exceed {len(FULL_LINE_STATIONS)}.")
    return FULL_LINE_STATIONS[:n_sections]


def sort_segments(values: Iterable[Any]) -> list[Any]:
    def key(value: Any) -> tuple[int, float | str]:
        try:
            return 0, float(value)
        except (TypeError, ValueError):
            return 1, str(value)

    return sorted(pd.Series(list(values)).dropna().unique().tolist(), key=key)


def first_valid(series: pd.Series) -> Any:
    values = series.dropna()
    return np.nan if values.empty else values.iloc[0]


def mode_or_first(series: pd.Series) -> Any:
    values = series.dropna()
    if values.empty:
        return np.nan
    counts = values.value_counts()
    return values.iloc[0] if counts.empty else counts.index[0]


def numeric_series(df: pd.DataFrame, col: str) -> pd.Series:
    if col not in df.columns:
        return pd.Series(dtype=float)
    return pd.to_numeric(df[col], errors="coerce")


def to_float(value: Any) -> float:
    if value is None or value == "":
        return math.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.nan


def mean(values: list[float]) -> float:
    return sum(values) / len(values) if values else math.nan


def percentile(values: list[float], q: float) -> float:
    if not values:
        return math.nan
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    pos = (len(ordered) - 1) * q
    lo = math.floor(pos)
    hi = math.ceil(pos)
    if lo == hi:
        return ordered[lo]
    return ordered[lo] * (hi - pos) + ordered[hi] * (pos - lo)


def cruise_average(values: list[float], min_speed_ratio: float) -> tuple[float, int]:
    if not values:
        return math.nan, 0
    max_speed = max(values)
    threshold = max_speed * min_speed_ratio
    cruise_values = [value for value in values if value >= threshold]
    return mean(cruise_values), len(cruise_values)


def format_float(value: Any) -> str:
    number = to_float(value)
    if not is_valid_number(number):
        return ""
    if float(number).is_integer():
        return str(int(number))
    return f"{number:.6f}".rstrip("0").rstrip(".")


def is_valid_number(value: float) -> bool:
    return not math.isnan(value)


def normalize_segment(value: Any) -> int | None:
    number = to_float(value)
    if not is_valid_number(number):
        return None
    rounded = round(number)
    if abs(number - rounded) <= 1e-9:
        return int(rounded)
    return None


def remember_first(store: dict[str, Any], key: str, value: Any) -> None:
    if key in store:
        return
    if value is None or value == "":
        return
    if isinstance(value, float) and math.isnan(value):
        return
    store[key] = value


def count_mode(store: dict[str, dict[Any, int]], key: str, value: Any) -> None:
    if value is None or value == "":
        return
    if isinstance(value, float) and math.isnan(value):
        return
    counts = store.setdefault(key, {})
    counts[value] = counts.get(value, 0) + 1


def get_mode(store: dict[str, dict[Any, int]], key: str) -> Any:
    counts = store.get(key, {})
    if not counts:
        return np.nan
    return max(counts.items(), key=lambda item: item[1])[0]


def new_aggregate(section: str, trip_no: int, segment_id: int, dwell: float, source_file: str) -> dict[str, Any]:
    return {
        "section": section,
        "trip_no": trip_no,
        "segment": segment_id,
        "dwell": dwell,
        "source_file": source_file,
        "count": 0,
        "min_time": math.inf,
        "max_time": -math.inf,
        "min_distance": math.inf,
        "max_distance": -math.inf,
        "energy_raw_sum": 0.0,
        "energy_seen": False,
        "speeds_kmh": [],
        "max_speed_kmh": -math.inf,
        "station_distance": math.nan,
        "mass_sum": 0.0,
        "mass_count": 0,
        "firsts": {},
        "modes": {},
    }


def aggregate_to_row(agg: dict[str, Any]) -> dict[str, Any]:
    if agg["count"] <= 0:
        raise ValueError(f"{agg['section']} trip {agg['trip_no']} has no rows.")

    start_time = agg["min_time"]
    end_time = agg["max_time"]
    run_time = end_time - start_time if math.isfinite(start_time) and math.isfinite(end_time) else math.nan

    distance = math.nan
    if math.isfinite(agg["min_distance"]) and math.isfinite(agg["max_distance"]):
        distance = agg["max_distance"] - agg["min_distance"]

    energy_wh = math.nan
    if agg["energy_seen"]:
        energy_wh = agg["energy_raw_sum"] / 3.6e6 * 1000.0

    avg_speed = distance / run_time * 3.6 if is_valid_number(distance) and is_valid_number(run_time) and run_time > 0 else math.nan
    speeds = agg["speeds_kmh"]
    sample_avg_speed = mean(speeds)
    cruise_avg_speed, cruise_point_count = cruise_average(speeds, agg["cruise_min_speed_ratio"])
    p95_speed = percentile(speeds, 0.95)
    p99_speed = percentile(speeds, 0.99)
    max_speed = agg["max_speed_kmh"] if math.isfinite(agg["max_speed_kmh"]) else math.nan
    avg_mass = agg["mass_sum"] / agg["mass_count"] if agg["mass_count"] > 0 else math.nan

    firsts: dict[str, Any] = agg["firsts"]
    dwell = float(agg["dwell"])
    return {
        SECTION_COL: agg["section"],
        "趟次": agg["trip_no"],
        "segment": agg["segment"],
        "日期+服务号": firsts.get("日期+服务号", np.nan),
        "服务号": firsts.get("服务号", np.nan),
        "车底号": firsts.get("车底号", np.nan),
        "方向": firsts.get("列车运行方向", np.nan),
        "时段": firsts.get("时段划分", np.nan),
        "开始时刻(s)": round(start_time, 2) if math.isfinite(start_time) else np.nan,
        "结束时刻(s)": round(end_time, 2) if math.isfinite(end_time) else np.nan,
        RUN_TIME_COL: round(run_time, 2) if is_valid_number(run_time) else np.nan,
        DWELL_COL: round(dwell, 2),
        TOTAL_TIME_COL: round(run_time + dwell, 2) if is_valid_number(run_time) else np.nan,
        DISTANCE_COL: round(distance, 2) if is_valid_number(distance) else np.nan,
        STATION_DISTANCE_COL: round(agg["station_distance"], 2) if is_valid_number(agg["station_distance"]) else np.nan,
        AVG_SPEED_COL: round(avg_speed, 2) if is_valid_number(avg_speed) else np.nan,
        MAX_SPEED_COL: round(max_speed, 2) if is_valid_number(max_speed) else np.nan,
        SAMPLE_AVG_SPEED_COL: round(sample_avg_speed, 2) if is_valid_number(sample_avg_speed) else np.nan,
        SPEED_SAMPLE_COUNT_COL: len(speeds),
        CRUISE_AVG_SPEED_COL: round(cruise_avg_speed, 2) if is_valid_number(cruise_avg_speed) else np.nan,
        CRUISE_POINT_COUNT_COL: cruise_point_count,
        P95_SPEED_COL: round(p95_speed, 2) if is_valid_number(p95_speed) else np.nan,
        P99_SPEED_COL: round(p99_speed, 2) if is_valid_number(p99_speed) else np.nan,
        ENERGY_COL: round(energy_wh, 2) if is_valid_number(energy_wh) else np.nan,
        MASS_COL: round(avg_mass, 2) if is_valid_number(avg_mass) else np.nan,
        "运行等级": get_mode(agg["modes"], "运行等级"),
        "是否晚上八点后": get_mode(agg["modes"], "是否是晚上八点以后的趟次"),
        "曲线质量标签": get_mode(agg["modes"], "曲线质量标签"),
        "源文件": agg["source_file"],
    }


def stream_trip_section_rows(
    data_dir: Path,
    section: str,
    trips: list[int],
    dwell_seconds: float,
    is_last_section: bool,
    cruise_min_speed_ratio: float,
) -> dict[int, dict[str, Any]]:
    """Read one section workbook once and summarize all requested trips.

    The Step2 files use 0-based `segment` values, so trip 1 maps to segment 0.
    Rows are normally sorted by segment. Once every requested segment has been
    seen and the stream moves past the largest target segment, reading stops.
    """

    from openpyxl import load_workbook

    file_path = data_dir / f"results_{section}.xlsx"
    if not file_path.exists():
        raise FileNotFoundError(f"Missing section file: {file_path}")

    source_file = str(file_path.relative_to(PROJECT_ROOT))
    target_by_segment = {trip_no - 1: trip_no for trip_no in trips}
    target_segments = set(target_by_segment)
    max_target_segment = max(target_segments)
    dwell = 0.0 if is_last_section else float(dwell_seconds)
    aggregates = {
        segment_id: new_aggregate(section, trip_no, segment_id, dwell, source_file)
        for segment_id, trip_no in target_by_segment.items()
    }
    for agg in aggregates.values():
        agg["cruise_min_speed_ratio"] = cruise_min_speed_ratio

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        except StopIteration as exc:
            raise ValueError(f"{file_path} is empty.") from exc

        col_index = {name: idx for idx, name in enumerate(headers)}
        if "segment" not in col_index:
            raise ValueError(f"{file_path} is missing required column: segment")
        if "时刻" not in col_index:
            raise ValueError(f"{file_path} is missing required column: 时刻")

        seen_segments: set[int] = set()
        for raw_row in rows:
            segment_id = normalize_segment(raw_row[col_index["segment"]])
            if segment_id is None:
                continue

            if seen_segments == target_segments and segment_id > max_target_segment:
                break
            if segment_id not in target_by_segment:
                continue

            seen_segments.add(segment_id)
            agg = aggregates[segment_id]
            agg["count"] += 1

            value = to_float(raw_row[col_index["时刻"]])
            if is_valid_number(value):
                agg["min_time"] = min(agg["min_time"], value)
                agg["max_time"] = max(agg["max_time"], value)

            if "累计位移(m)" in col_index:
                value = to_float(raw_row[col_index["累计位移(m)"]])
                if is_valid_number(value):
                    agg["min_distance"] = min(agg["min_distance"], value)
                    agg["max_distance"] = max(agg["max_distance"], value)

            if "energy" in col_index:
                value = to_float(raw_row[col_index["energy"]])
                if is_valid_number(value):
                    agg["energy_raw_sum"] += value
                    agg["energy_seen"] = True

            if "速度(m/s)" in col_index:
                value = to_float(raw_row[col_index["速度(m/s)"]])
                if is_valid_number(value) and value >= 0:
                    speed_kmh = value * 3.6
                    agg["speeds_kmh"].append(speed_kmh)
                    agg["max_speed_kmh"] = max(agg["max_speed_kmh"], speed_kmh)

            if "站间距" in col_index and not is_valid_number(agg["station_distance"]):
                value = to_float(raw_row[col_index["站间距"]])
                if is_valid_number(value):
                    agg["station_distance"] = value

            if "重量" in col_index:
                value = to_float(raw_row[col_index["重量"]])
                if is_valid_number(value):
                    agg["mass_sum"] += value
                    agg["mass_count"] += 1

            for name in ["日期+服务号", "服务号", "车底号", "列车运行方向", "时段划分"]:
                if name in col_index:
                    remember_first(agg["firsts"], name, raw_row[col_index[name]])

            for name in ["运行等级", "是否是晚上八点以后的趟次", "曲线质量标签"]:
                if name in col_index:
                    count_mode(agg["modes"], name, raw_row[col_index[name]])
    finally:
        wb.close()

    rows_by_trip: dict[int, dict[str, Any]] = {}
    for segment_id, trip_no in target_by_segment.items():
        agg = aggregates[segment_id]
        if agg["count"] > 0:
            rows_by_trip[trip_no] = aggregate_to_row(agg)
    return rows_by_trip


def read_section_file(data_dir: Path, section: str) -> pd.DataFrame:
    file_path = data_dir / f"results_{section}.xlsx"
    if not file_path.exists():
        raise FileNotFoundError(f"Missing section file: {file_path}")

    header = pd.read_excel(file_path, nrows=0)
    usecols = [col for col in SOURCE_COLS if col in header.columns]
    if "segment" not in usecols:
        raise ValueError(f"{file_path} is missing required column: segment")
    if "时刻" not in usecols:
        raise ValueError(f"{file_path} is missing required column: 时刻")

    df = pd.read_excel(file_path, usecols=usecols)
    df["__source_file"] = str(file_path.relative_to(PROJECT_ROOT))
    return df


def summarize_trip_section(
    section: str,
    df_all: pd.DataFrame,
    trip_no: int,
    dwell_seconds: float,
    is_last_section: bool,
) -> dict[str, Any]:
    trip_index = trip_no - 1
    segments = sort_segments(df_all["segment"])
    if trip_index >= len(segments):
        raise IndexError(f"{section} only has {len(segments)} segment(s); trip {trip_no} is unavailable.")

    segment_id = segments[trip_index]
    df = df_all[df_all["segment"] == segment_id].copy()
    if df.empty:
        raise ValueError(f"{section} trip {trip_no} has no rows.")

    df["时刻"] = numeric_series(df, "时刻")
    df = df.sort_values("时刻")
    time_values = df["时刻"].dropna()
    if time_values.empty:
        raise ValueError(f"{section} trip {trip_no} has no valid 时刻 values.")

    start_time = float(time_values.iloc[0])
    end_time = float(time_values.iloc[-1])
    run_time = end_time - start_time

    distance = np.nan
    displacement = numeric_series(df, "累计位移(m)").dropna()
    if not displacement.empty:
        distance = float(displacement.max() - displacement.min())

    energy_wh = np.nan
    energy = numeric_series(df, "energy").dropna()
    if not energy.empty:
        energy_wh = float(energy.sum() / 3.6e6 * 1000.0)

    speed_kmh = numeric_series(df, "速度(m/s)") * 3.6
    valid_speed = speed_kmh.dropna()
    max_speed = float(valid_speed.max()) if not valid_speed.empty else np.nan
    avg_speed = distance / run_time * 3.6 if pd.notna(distance) and run_time > 0 else np.nan

    station_distance = first_valid(numeric_series(df, "站间距")) if "站间距" in df.columns else np.nan
    mass_values = numeric_series(df, "重量").dropna()
    avg_mass = float(mass_values.mean()) if not mass_values.empty else np.nan

    dwell = 0.0 if is_last_section else float(dwell_seconds)

    return {
        SECTION_COL: section,
        "趟次": trip_no,
        "segment": segment_id,
        "日期+服务号": first_valid(df["日期+服务号"]) if "日期+服务号" in df.columns else np.nan,
        "服务号": first_valid(df["服务号"]) if "服务号" in df.columns else np.nan,
        "车底号": first_valid(df["车底号"]) if "车底号" in df.columns else np.nan,
        "方向": first_valid(df["列车运行方向"]) if "列车运行方向" in df.columns else np.nan,
        "时段": first_valid(df["时段划分"]) if "时段划分" in df.columns else np.nan,
        "开始时刻(s)": round(start_time, 2),
        "结束时刻(s)": round(end_time, 2),
        RUN_TIME_COL: round(run_time, 2),
        DWELL_COL: round(dwell, 2),
        TOTAL_TIME_COL: round(run_time + dwell, 2),
        DISTANCE_COL: round(distance, 2) if pd.notna(distance) else np.nan,
        STATION_DISTANCE_COL: round(float(station_distance), 2) if pd.notna(station_distance) else np.nan,
        AVG_SPEED_COL: round(avg_speed, 2) if pd.notna(avg_speed) else np.nan,
        MAX_SPEED_COL: round(max_speed, 2) if pd.notna(max_speed) else np.nan,
        ENERGY_COL: round(energy_wh, 2) if pd.notna(energy_wh) else np.nan,
        MASS_COL: round(avg_mass, 2) if pd.notna(avg_mass) else np.nan,
        "运行等级": mode_or_first(df["运行等级"]) if "运行等级" in df.columns else np.nan,
        "是否晚上八点后": mode_or_first(df["是否是晚上八点以后的趟次"]) if "是否是晚上八点以后的趟次" in df.columns else np.nan,
        "曲线质量标签": mode_or_first(df["曲线质量标签"]) if "曲线质量标签" in df.columns else np.nan,
        "源文件": first_valid(df["__source_file"]),
    }


def append_total_row(df: pd.DataFrame) -> pd.DataFrame:
    total_run = pd.to_numeric(df[RUN_TIME_COL], errors="coerce").sum(skipna=True)
    total_dwell = pd.to_numeric(df[DWELL_COL], errors="coerce").sum(skipna=True)
    total_distance = pd.to_numeric(df[DISTANCE_COL], errors="coerce").sum(skipna=True)
    total_station_distance = pd.to_numeric(df[STATION_DISTANCE_COL], errors="coerce").sum(skipna=True)
    total_energy = pd.to_numeric(df[ENERGY_COL], errors="coerce").sum(skipna=True)
    max_speed = pd.to_numeric(df[MAX_SPEED_COL], errors="coerce").max(skipna=True)
    avg_speed = total_distance / total_run * 3.6 if total_run > 0 else np.nan

    total_row = {col: "" for col in df.columns}
    total_row.update(
        {
            SECTION_COL: "--- 总计 ---",
            RUN_TIME_COL: round(total_run, 2),
            DWELL_COL: round(total_dwell, 2),
            TOTAL_TIME_COL: round(total_run + total_dwell, 2),
            DISTANCE_COL: round(total_distance, 2),
            STATION_DISTANCE_COL: round(total_station_distance, 2),
            AVG_SPEED_COL: round(avg_speed, 2) if pd.notna(avg_speed) else np.nan,
            MAX_SPEED_COL: round(max_speed, 2) if pd.notna(max_speed) else np.nan,
            SAMPLE_AVG_SPEED_COL: "",
            SPEED_SAMPLE_COUNT_COL: "",
            CRUISE_AVG_SPEED_COL: "",
            CRUISE_POINT_COUNT_COL: "",
            P95_SPEED_COL: "",
            P99_SPEED_COL: "",
            ENERGY_COL: round(total_energy, 2),
            MASS_COL: "...",
        }
    )
    return pd.concat([df, pd.DataFrame([total_row])], ignore_index=True)


def report_row_to_lookup(section_index: int, row: dict[str, Any], cruise_min_speed_ratio: float) -> dict[str, Any]:
    return {
        "section_index": section_index,
        "section": row.get(SECTION_COL, ""),
        "segment": row.get("segment", ""),
        "trip_no": row.get("趟次", ""),
        "run_id": row.get("日期+服务号", ""),
        "quality_labels": row.get("曲线质量标签", ""),
        "row_count": row.get(SPEED_SAMPLE_COUNT_COL, ""),
        "speed_source": "mps",
        "avg_speed_kmh": format_float(row.get(SAMPLE_AVG_SPEED_COL)),
        "cruise_avg_speed_kmh": format_float(row.get(CRUISE_AVG_SPEED_COL)),
        "cruise_point_count": row.get(CRUISE_POINT_COUNT_COL, ""),
        "cruise_min_speed_ratio": format_float(cruise_min_speed_ratio),
        "max_speed_kmh": format_float(row.get(MAX_SPEED_COL)),
        "p99_speed_kmh": format_float(row.get(P99_SPEED_COL)),
        "p95_speed_kmh": format_float(row.get(P95_SPEED_COL)),
        "xlsx_file": str(PROJECT_ROOT / str(row.get("源文件", ""))) if row.get("源文件") else "",
        "warning": "",
    }


def write_trip_report(trip_dir: Path, trip_no: int, rows: list[dict[str, Any]], output_name: str) -> dict[str, Any]:
    trip_dir.mkdir(parents=True, exist_ok=True)
    df = append_total_row(pd.DataFrame(rows))

    csv_path = trip_dir / f"{output_name}.csv"
    xlsx_path = trip_dir / f"{output_name}.xlsx"
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="历史运行情况")
        ws = writer.book["历史运行情况"]
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 36)

    total = df.iloc[-1].to_dict()
    return {
        "trip_no": trip_no,
        "trip_label": f"trip{trip_no:03d}",
        "section_count": len(rows),
        "historical_run_time_s": total[RUN_TIME_COL],
        "historical_dwell_time_s": total[DWELL_COL],
        "historical_total_time_s": total[TOTAL_TIME_COL],
        "historical_energy_wh": total[ENERGY_COL],
        "csv": str(csv_path),
        "xlsx": str(xlsx_path),
    }


def write_speed_lookup_outputs(
    output_dir: Path,
    lookup_rows: list[dict[str, Any]],
    output_name: str,
) -> None:
    if not lookup_rows:
        return

    all_lookup_path = output_dir / "historical_speed_lookup.csv"
    write_csv_rows(all_lookup_path, lookup_rows)
    print(f"Speed lookup: {all_lookup_path}")

    rows_by_trip: dict[int, list[dict[str, Any]]] = {}
    for row in lookup_rows:
        trip_no = int(float(row["trip_no"]))
        rows_by_trip.setdefault(trip_no, []).append(row)

    for trip_no, rows in rows_by_trip.items():
        trip_path = output_dir / f"trip{trip_no:03d}" / f"{output_name}_speed_lookup.csv"
        write_csv_rows(trip_path, rows)


def write_csv_rows(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def write_batch_summary(output_dir: Path, rows: list[dict[str, Any]], failures: list[dict[str, Any]]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    summary_path = output_dir / "historical_only_batch_summary.csv"
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    if fieldnames:
        write_csv_rows(summary_path, rows)
        print(f"Batch summary: {summary_path}")

    if failures:
        failure_path = output_dir / "historical_only_failures.csv"
        write_csv_rows(failure_path, failures)
        print(f"Failure summary: {failure_path}")


def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Batch-build historical-only per-trip section reports.")
    parser.add_argument("--trips", required=True, help="Trips to extract, e.g. 1,2,5-10.")
    parser.add_argument("--line-scope", default="full", help="'full', 'all', or first N sections, e.g. 5.")
    parser.add_argument("--sections", help="Comma-separated section names. Overrides --line-scope.")
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Directory containing results_<section>.xlsx files.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Directory containing/receiving tripXXX folders.")
    parser.add_argument("--output-name", default="Historical_Only_Report", help="Base file name written inside each trip folder.")
    parser.add_argument("--history-dwell", type=float, default=30.0, help="Historical dwell seconds between sections.")
    parser.add_argument(
        "--cruise-min-speed-ratio",
        type=float,
        default=0.92,
        help="Speed samples >= max_speed * this ratio are averaged as cruise speed for later OpenTrack speed limits.",
    )
    parser.add_argument(
        "--history-dwell-mode",
        choices=["fixed", "zero"],
        default="fixed",
        help="fixed uses --history-dwell for all but the last selected section; zero ignores dwell.",
    )
    parser.add_argument("--continue-on-error", action="store_true", help="Continue if a section or trip fails.")
    return parser


def main() -> int:
    args = build_arg_parser().parse_args()
    trips = parse_trip_list(args.trips)
    sections = choose_sections(args.line_scope, args.sections)
    data_dir = resolve_project_path(args.data_dir)
    output_dir = resolve_project_path(args.output_dir)
    dwell_seconds = 0.0 if args.history_dwell_mode == "zero" else args.history_dwell

    if not data_dir.exists():
        raise FileNotFoundError(f"Data directory does not exist: {data_dir}")

    print(f"Data dir:   {data_dir}")
    print(f"Output dir: {output_dir}")
    print(f"Trips:      {trips}")
    print(f"Sections:   {len(sections)}")
    print(f"Dwell:      {dwell_seconds:g}s ({args.history_dwell_mode})")

    rows_by_trip: dict[int, list[dict[str, Any]]] = {trip: [] for trip in trips}
    lookup_rows: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []

    for section_index, section in enumerate(sections, 1):
        print(f"[{section_index}/{len(sections)}] Reading {section}...")
        try:
            section_rows = stream_trip_section_rows(
                data_dir=data_dir,
                section=section,
                trips=trips,
                dwell_seconds=dwell_seconds,
                is_last_section=(section_index == len(sections)),
                cruise_min_speed_ratio=args.cruise_min_speed_ratio,
            )
        except Exception as exc:
            message = str(exc)
            print(f"  ERROR read section: {message}")
            failures.append({"section": section, "trip_no": "", "error": message})
            if not args.continue_on_error:
                raise
            continue

        for trip_no in trips:
            try:
                if trip_no not in section_rows:
                    raise ValueError(f"{section} trip {trip_no} has no matching segment rows.")
                row = section_rows[trip_no]
                rows_by_trip[trip_no].append(row)
                lookup_rows.append(report_row_to_lookup(section_index, row, args.cruise_min_speed_ratio))
            except Exception as exc:
                message = str(exc)
                print(f"  ERROR trip {trip_no}: {message}")
                failures.append({"section": section, "trip_no": trip_no, "error": message})
                if not args.continue_on_error:
                    raise

    summary_rows: list[dict[str, Any]] = []
    for trip_no in trips:
        trip_rows = rows_by_trip[trip_no]
        if not trip_rows:
            failures.append({"section": "", "trip_no": trip_no, "error": "No rows generated for trip."})
            if not args.continue_on_error:
                raise RuntimeError(f"No rows generated for trip {trip_no}.")
            continue

        trip_dir = output_dir / f"trip{trip_no:03d}"
        print(f"Writing trip{trip_no:03d}: {len(trip_rows)} section(s)")
        summary_rows.append(write_trip_report(trip_dir, trip_no, trip_rows, args.output_name))

    write_speed_lookup_outputs(output_dir, lookup_rows, args.output_name)
    write_batch_summary(output_dir, summary_rows, failures)

    if failures:
        print(f"Finished with {len(failures)} warning/error record(s).")
        return 1 if not summary_rows else 0
    print("Finished successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
