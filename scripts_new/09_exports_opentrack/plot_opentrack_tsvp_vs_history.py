from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import openpyxl


PROJECT_ROOT = Path(r"D:\energy_conservation")
DEFAULT_TSVP = Path(r"D:\OutPut\OT_priority_history_trip001.tsvP")
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_all_curve_quality"
DEFAULT_ROUTE_MAP = PROJECT_ROOT / "output" / "opentrack_route_map_newdoc" / "priority_dp_first5_route_map.csv"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output" / "opentrack_tsvp_compare"

SECTION_COL = "section"
TIME_COL = "\u65f6\u523b"
SPEED_COL = "\u901f\u5ea6(m/s)"
DIST_COL = "\u7d2f\u8ba1\u4f4d\u79fb(m)"
RUN_ID_COL = "\u65e5\u671f+\u670d\u52a1\u53f7"
QUALITY_COL = "\u66f2\u7ebf\u8d28\u91cf\u6807\u7b7e"
SEGMENT_COL = "segment"
STEP_ENERGY_COL = "energy"
CUM_ENERGY_COL = "cumulative_energy"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Plot OpenTrack .tsvP curves against Step2 historical curves: "
            "v-t, v-s, E-t and E-s."
        )
    )
    parser.add_argument("--tsvp", default=str(DEFAULT_TSVP), help="OpenTrack .tsvP output file.")
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Step2 results_*.xlsx directory.")
    parser.add_argument("--route-map", default=str(DEFAULT_ROUTE_MAP), help="Route map CSV for section order.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Directory for plots and summary CSV.")
    parser.add_argument("--prefix", default=None, help="Output filename prefix. Defaults to the .tsvP stem.")
    parser.add_argument("--trip-no", type=int, default=1, help="Historical trip number, 1-based.")
    parser.add_argument("--run-id", default=None, help="Exact historical run id. Overrides --trip-no.")
    parser.add_argument("--section-start", type=int, default=1, help="First 1-based section index to include.")
    parser.add_argument("--section-count", type=int, default=None, help="Only include this many sections.")
    parser.add_argument(
        "--sections",
        default=None,
        help="Optional comma-separated section names. Overrides --section-start/--section-count.",
    )
    parser.add_argument(
        "--history-dwell",
        type=float,
        default=0.0,
        help="Optional dwell seconds inserted between historical sections for timeline comparison.",
    )
    parser.add_argument(
        "--quality-label",
        default=None,
        help="Optional historical curve quality filter, for example 0 for normal curves only.",
    )
    parser.add_argument("--dpi", type=int, default=300, help="Output image DPI.")
    parser.add_argument("--image-format", choices=["png", "jpg"], default="png", help="Output image format.")
    parser.add_argument("--show", action="store_true", help="Show the plot window after saving.")
    return parser.parse_args()


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def selected_sections(route_map: Path, args: argparse.Namespace) -> list[str]:
    rows = read_csv(route_map)
    sections = [(row.get(SECTION_COL) or "").strip() for row in rows]
    sections = [section for section in sections if section]
    if args.sections:
        wanted = {item.strip() for item in args.sections.split(",") if item.strip()}
        return [section for section in sections if section in wanted]
    start = max(args.section_start, 1) - 1
    end = None if args.section_count is None else start + args.section_count
    return sections[start:end]


def read_tsvp(path: Path) -> dict[str, list[float]]:
    time_abs: list[float] = []
    dist_km: list[float] = []
    speed_kmh: list[float] = []
    energy_kwh: list[float] = []

    with path.open("r", encoding="utf-8", errors="replace") as f:
        for line in f:
            if line.startswith("//") or not line.strip():
                continue
            parts = [part.strip() for part in line.rstrip("\n").split("\t")]
            parts = [part for part in parts if part != ""]
            if len(parts) < 9:
                continue
            t = as_float(parts[0])
            s = as_float(parts[1])
            v = as_float(parts[2])
            e_mj = as_float(parts[8])
            if t is None or s is None or v is None or e_mj is None:
                continue
            time_abs.append(t)
            dist_km.append(s)
            speed_kmh.append(v)
            energy_kwh.append(e_mj / 3.6)

    if not time_abs:
        raise ValueError(f"No numeric .tsvP rows found: {path}")

    t0 = time_abs[0]
    s0 = dist_km[0]
    e0 = energy_kwh[0]
    return {
        "time_s": [t - t0 for t in time_abs],
        "distance_km": [s - s0 for s in dist_km],
        "speed_kmh": speed_kmh,
        "energy_kwh": [e - e0 for e in energy_kwh],
    }


def header_index(headers: list[str], name: str) -> int | None:
    normalized = {str(value).strip(): idx for idx, value in enumerate(headers)}
    return normalized.get(name)


def records_to_history_curve(
    records: list[tuple[float, float, float, float | None, float | None]],
) -> dict[str, list[float]]:
    records.sort(key=lambda item: item[0])
    t0 = records[0][0]
    s0 = records[0][2]
    first_cum = records[0][4]

    time_s: list[float] = []
    speed_kmh: list[float] = []
    distance_km: list[float] = []
    energy_kwh: list[float] = []
    running_energy_j = 0.0

    for t, v, s, e_step, e_cum in records:
        time_s.append(t - t0)
        speed_kmh.append(v)
        distance_km.append(s - s0)
        if e_cum is not None and first_cum is not None:
            energy_kwh.append((e_cum - first_cum) / 3_600_000.0)
        else:
            running_energy_j += e_step or 0.0
            energy_kwh.append(running_energy_j / 3_600_000.0)

    return {
        "time_s": time_s,
        "speed_kmh": speed_kmh,
        "distance_km": distance_km,
        "energy_kwh": energy_kwh,
    }


def read_cached_history_section(
    cache_path: Path,
    section: str,
    args: argparse.Namespace,
) -> dict[str, list[float]]:
    import pandas as pd

    frame = pd.read_pickle(cache_path)
    required = [TIME_COL, SPEED_COL, DIST_COL]
    missing = [column for column in required if column not in frame.columns]
    if missing:
        raise ValueError(f"{cache_path.name}: missing cached columns {missing}")
    if CUM_ENERGY_COL not in frame.columns and STEP_ENERGY_COL not in frame.columns:
        raise ValueError(f"{cache_path.name}: missing cached energy/cumulative_energy column")

    if args.run_id and RUN_ID_COL in frame.columns:
        frame = frame[frame[RUN_ID_COL].astype(str).str.strip().eq(str(args.run_id).strip())]
    if args.quality_label is not None and QUALITY_COL in frame.columns:
        frame = frame[
            frame[QUALITY_COL].astype(str).str.strip().eq(str(args.quality_label).strip())
        ]

    records: list[tuple[float, float, float, float | None, float | None]] = []
    has_step_energy = STEP_ENERGY_COL in frame.columns
    has_cum_energy = CUM_ENERGY_COL in frame.columns
    for _, row in frame.iterrows():
        t = as_float(row.get(TIME_COL))
        v_ms = as_float(row.get(SPEED_COL))
        s_m = as_float(row.get(DIST_COL))
        e_step = as_float(row.get(STEP_ENERGY_COL)) if has_step_energy else None
        e_cum = as_float(row.get(CUM_ENERGY_COL)) if has_cum_energy else None
        if t is None or v_ms is None or s_m is None:
            continue
        records.append((t, v_ms * 3.6, s_m / 1000.0, e_step, e_cum))

    if not records:
        raise ValueError(f"{cache_path.name}: no cached historical rows selected for {section}")
    return records_to_history_curve(records)


def read_history_section(
    xlsx_path: Path,
    section: str,
    args: argparse.Namespace,
) -> dict[str, list[float]]:
    cache_by_section = getattr(args, "history_section_cache", None) or {}
    cache_path = cache_by_section.get(section)
    if cache_path and Path(cache_path).is_file():
        return read_cached_history_section(Path(cache_path), section, args)

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Missing Step2 xlsx for {section}: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]

        idx_time = header_index(headers, TIME_COL)
        idx_speed = header_index(headers, SPEED_COL)
        idx_dist = header_index(headers, DIST_COL)
        idx_segment = header_index(headers, SEGMENT_COL)
        idx_run_id = header_index(headers, RUN_ID_COL)
        idx_quality = header_index(headers, QUALITY_COL)
        idx_energy = header_index(headers, STEP_ENERGY_COL)
        idx_cum_energy = header_index(headers, CUM_ENERGY_COL)

        required = {
            "time": idx_time,
            "speed": idx_speed,
            "distance": idx_dist,
            "segment": idx_segment,
        }
        missing = [name for name, idx in required.items() if idx is None]
        if missing:
            raise ValueError(f"{xlsx_path.name}: missing columns {missing}")
        if idx_cum_energy is None and idx_energy is None:
            raise ValueError(f"{xlsx_path.name}: missing energy/cumulative_energy column")

        segment_by_section = getattr(args, "segment_by_section", None) or {}
        target_segment = int(segment_by_section.get(section, args.trip_no - 1))
        records: list[tuple[float, float, float, float | None, float | None]] = []
        seen_target_segment = False

        for row in rows:
            if args.run_id and idx_run_id is not None:
                current_run_id = "" if row[idx_run_id] is None else str(row[idx_run_id]).strip()
                if current_run_id != args.run_id:
                    continue
            else:
                segment = as_float(row[idx_segment]) if idx_segment is not None else None
                if segment is None:
                    continue
                current_segment = int(segment)
                if seen_target_segment and current_segment != target_segment:
                    break
                if current_segment != target_segment:
                    continue
                seen_target_segment = True

            if args.quality_label is not None and idx_quality is not None:
                quality = "" if row[idx_quality] is None else str(row[idx_quality]).strip()
                if quality != str(args.quality_label):
                    continue

            t = as_float(row[idx_time])
            v_ms = as_float(row[idx_speed])
            s_m = as_float(row[idx_dist])
            e_step = as_float(row[idx_energy]) if idx_energy is not None else None
            e_cum = as_float(row[idx_cum_energy]) if idx_cum_energy is not None else None
            if t is None or v_ms is None or s_m is None:
                continue
            records.append((t, v_ms * 3.6, s_m / 1000.0, e_step, e_cum))

        if not records:
            selector = f"run_id={args.run_id}" if args.run_id else f"trip_no={args.trip_no}"
            raise ValueError(f"{xlsx_path.name}: no rows selected for {selector}")
        return records_to_history_curve(records)
    finally:
        wb.close()


def append_history_section(
    full: dict[str, list[float]],
    section_data: dict[str, list[float]],
    dwell_s: float,
    is_first: bool,
) -> None:
    if not section_data["time_s"]:
        return

    if full["time_s"]:
        t_offset = full["time_s"][-1]
        s_offset = full["distance_km"][-1]
        e_offset = full["energy_kwh"][-1]
        if not is_first and dwell_s > 0:
            full["time_s"].append(t_offset + dwell_s)
            full["distance_km"].append(s_offset)
            full["speed_kmh"].append(0.0)
            full["energy_kwh"].append(e_offset)
            t_offset = full["time_s"][-1]
    else:
        t_offset = 0.0
        s_offset = 0.0
        e_offset = 0.0

    start_idx = 0 if is_first else 1
    for idx in range(start_idx, len(section_data["time_s"])):
        full["time_s"].append(t_offset + section_data["time_s"][idx])
        full["distance_km"].append(s_offset + section_data["distance_km"][idx])
        full["speed_kmh"].append(section_data["speed_kmh"][idx])
        full["energy_kwh"].append(e_offset + section_data["energy_kwh"][idx])


def read_history_line(sections: list[str], args: argparse.Namespace) -> tuple[dict[str, list[float]], list[dict[str, Any]]]:
    full = {"time_s": [], "distance_km": [], "speed_kmh": [], "energy_kwh": []}
    summary_rows: list[dict[str, Any]] = []

    for idx, section in enumerate(sections):
        xlsx_path = Path(args.data_dir) / f"results_{section}.xlsx"
        section_data = read_history_section(xlsx_path, section, args)
        run_time_by_section = getattr(args, "history_run_time_by_section", None) or {}
        target_run_time = run_time_by_section.get(section)
        if target_run_time is not None and section_data["time_s"]:
            raw_run_time = section_data["time_s"][-1]
            if raw_run_time > 0:
                scale = float(target_run_time) / raw_run_time
                section_data["time_s"] = [value * scale for value in section_data["time_s"]]
        dwell_by_section = getattr(args, "history_dwell_by_section", None) or {}
        dwell_key = sections[idx - 1] if idx > 0 else section
        dwell_s = float(dwell_by_section.get(dwell_key, args.history_dwell))
        append_history_section(full, section_data, dwell_s, idx == 0)
        summary_rows.append(
            {
                "source": "history",
                "section_index": idx + 1,
                "section": section,
                "points": len(section_data["time_s"]),
                "duration_s": section_data["time_s"][-1] if section_data["time_s"] else "",
                "distance_km": section_data["distance_km"][-1] if section_data["distance_km"] else "",
                "energy_kwh": section_data["energy_kwh"][-1] if section_data["energy_kwh"] else "",
            }
        )
    return full, summary_rows


def curve_summary(source: str, curve: dict[str, list[float]]) -> dict[str, Any]:
    return {
        "source": source,
        "section_index": "",
        "section": "TOTAL",
        "points": len(curve["time_s"]),
        "duration_s": curve["time_s"][-1] if curve["time_s"] else "",
        "distance_km": curve["distance_km"][-1] if curve["distance_km"] else "",
        "energy_kwh": curve["energy_kwh"][-1] if curve["energy_kwh"] else "",
    }


def energy_error(ot: dict[str, list[float]], history: dict[str, list[float]]) -> tuple[float | None, float | None]:
    if not ot["energy_kwh"] or not history["energy_kwh"]:
        return None, None
    ot_final = ot["energy_kwh"][-1]
    history_final = history["energy_kwh"][-1]
    diff = ot_final - history_final
    if abs(history_final) < 1e-9:
        return diff, None
    return diff, diff / history_final * 100.0


def add_energy_error_box(ax: Any, diff_kwh: float | None, pct: float | None) -> None:
    if diff_kwh is None:
        text = "\u80fd\u8017\u8bef\u5dee: N/A"
    elif pct is None:
        text = f"\u80fd\u8017\u8bef\u5dee: N/A\nOT-History: {diff_kwh:+.3f} kWh"
    else:
        text = f"\u80fd\u8017\u8bef\u5dee: {pct:+.2f}%\nOT-History: {diff_kwh:+.3f} kWh"

    ax.text(
        0.98,
        0.05,
        text,
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=10,
        bbox={
            "boxstyle": "round,pad=0.35",
            "facecolor": "white",
            "edgecolor": "#999999",
            "alpha": 0.9,
        },
    )


def plot_compare(
    ot: dict[str, list[float]],
    history: dict[str, list[float]],
    output_png: Path,
    title: str,
    dpi: int,
    show: bool,
) -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["lines.antialiased"] = True

    fig, axes = plt.subplots(2, 2, figsize=(16, 10))
    ax_vt, ax_vs, ax_et, ax_es = axes.ravel()
    diff_kwh, diff_pct = energy_error(ot, history)

    ax_vt.plot(ot["time_s"], ot["speed_kmh"], color="#1f77b4", lw=2.0, label="OpenTrack")
    ax_vt.plot(history["time_s"], history["speed_kmh"], color="#111111", lw=1.8, ls="--", label="History")
    ax_vt.set_title("v-t")
    ax_vt.set_xlabel("Time (s)")
    ax_vt.set_ylabel("Speed (km/h)")

    ax_vs.plot(ot["distance_km"], ot["speed_kmh"], color="#1f77b4", lw=2.0, label="OpenTrack")
    ax_vs.plot(history["distance_km"], history["speed_kmh"], color="#111111", lw=1.8, ls="--", label="History")
    ax_vs.set_title("v-s")
    ax_vs.set_xlabel("Distance (km)")
    ax_vs.set_ylabel("Speed (km/h)")

    ax_et.plot(ot["time_s"], ot["energy_kwh"], color="#d62728", lw=2.0, label="OpenTrack")
    ax_et.plot(history["time_s"], history["energy_kwh"], color="#111111", lw=1.8, ls="--", label="History")
    ax_et.set_title("E-t")
    ax_et.set_xlabel("Time (s)")
    ax_et.set_ylabel("Cumulative Energy (kWh)")
    add_energy_error_box(ax_et, diff_kwh, diff_pct)

    ax_es.plot(ot["distance_km"], ot["energy_kwh"], color="#d62728", lw=2.0, label="OpenTrack")
    ax_es.plot(history["distance_km"], history["energy_kwh"], color="#111111", lw=1.8, ls="--", label="History")
    ax_es.set_title("E-s")
    ax_es.set_xlabel("Distance (km)")
    ax_es.set_ylabel("Cumulative Energy (kWh)")
    add_energy_error_box(ax_es, diff_kwh, diff_pct)

    for ax in axes.ravel():
        ax.grid(True, alpha=0.25)
        ax.legend(loc="best")

    fig.suptitle(title, fontsize=15)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    output_png.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_png, dpi=dpi, bbox_inches="tight")
    if show:
        plt.show()
    plt.close(fig)


def main() -> int:
    args = parse_args()
    tsvp_path = Path(args.tsvp)
    route_map = Path(args.route_map)
    sections = selected_sections(route_map, args)
    if not sections:
        raise ValueError("No sections selected.")

    print(f"Reading OpenTrack .tsvP: {tsvp_path}")
    ot_curve = read_tsvp(tsvp_path)
    print(f"Reading history sections: {len(sections)}")
    history_curve, history_summary = read_history_line(sections, args)

    prefix = args.prefix or tsvp_path.stem
    output_dir = Path(args.output_dir)
    selector = f"run_id={args.run_id}" if args.run_id else f"trip_no={args.trip_no}"
    output_png = output_dir / f"{prefix}_vs_history_{selector.replace('/', '-')}_vt_vs_et_es.{args.image_format}"
    summary_csv = output_dir / f"{prefix}_vs_history_{selector.replace('/', '-')}_summary.csv"

    title = (
        f"{tsvp_path.stem} vs history | {selector} | "
        f"{len(sections)} section(s) | history dwell={args.history_dwell:g}s"
    )
    plot_compare(ot_curve, history_curve, output_png, title, args.dpi, args.show)

    diff_kwh, diff_pct = energy_error(ot_curve, history_curve)
    summary_rows = [curve_summary("opentrack", ot_curve), curve_summary("history_total", history_curve)]
    summary_rows.append(
        {
            "source": "energy_error",
            "section_index": "",
            "section": "OpenTrack - History",
            "points": "",
            "duration_s": "",
            "distance_km": "",
            "energy_kwh": diff_kwh if diff_kwh is not None else "",
            "energy_error_pct": diff_pct if diff_pct is not None else "",
        }
    )
    summary_rows.extend(history_summary)
    write_csv(summary_csv, summary_rows)

    print(f"Sections:    {len(sections)}")
    print(f"OpenTrack:   {ot_curve['time_s'][-1]:.1f}s, {ot_curve['distance_km'][-1]:.3f} km, {ot_curve['energy_kwh'][-1]:.3f} kWh")
    print(
        f"History:     {history_curve['time_s'][-1]:.1f}s, "
        f"{history_curve['distance_km'][-1]:.3f} km, {history_curve['energy_kwh'][-1]:.3f} kWh"
    )
    if diff_kwh is not None and diff_pct is not None:
        print(f"Energy err:  {diff_pct:+.2f}% ({diff_kwh:+.3f} kWh, OpenTrack - History)")
    print(f"Plot:        {output_png}")
    print(f"Summary CSV: {summary_csv}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
