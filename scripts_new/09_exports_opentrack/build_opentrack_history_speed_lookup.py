from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path
from typing import Any

import openpyxl


PROJECT_ROOT = Path(r"D:\energy_conservation")
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_all_curve_quality"
DEFAULT_ROUTE_MAP = PROJECT_ROOT / "output" / "opentrack_route_map" / "priority_history_probe_route_map.csv"
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "opentrack_speed_limits" / "history_speed_lookup.csv"

# Keep Chinese column names as unicode escapes so this script is stable under
# different Windows console code pages.
SECTION_COLS = ["\u533a\u6bb5", "\u7ad9\u95f4\u533a\u95f4", "section"]
RUN_ID_COLS = ["\u65e5\u671f+\u670d\u52a1\u53f7"]
SEGMENT_COLS = ["segment"]
QUALITY_COLS = ["\u66f2\u7ebf\u8d28\u91cf\u6807\u7b7e"]
REFERENCE_SPEED_COLS = ["BCU6_CCU \u53c2\u8003\u901f\u5ea6km/h", "\u53c2\u8003\u901f\u5ea6km/h"]
MPS_SPEED_COLS = ["\u901f\u5ea6(m/s)"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Build a reusable historical speed lookup table from Step2 processed xlsx files. "
            "The output can be used later without scanning the large xlsx files again."
        )
    )
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Step2 processed xlsx directory.")
    parser.add_argument(
        "--route-map",
        default=None,
        help=(
            "Optional route map CSV. When provided, its section order is used and "
            "--section-start/--section-count can select a contiguous subset."
        ),
    )
    parser.add_argument("--output-csv", default=str(DEFAULT_OUTPUT), help="Lookup CSV to write.")
    parser.add_argument("--section-start", type=int, default=1, help="First 1-based section index to include.")
    parser.add_argument("--section-count", type=int, default=None, help="Only include this many sections.")
    parser.add_argument(
        "--sections",
        default=None,
        help="Optional comma-separated section names to include, overriding --section-start/--section-count.",
    )
    parser.add_argument(
        "--speed-source",
        choices=["reference", "mps", "auto"],
        default="auto",
        help="Speed source: reference speed km/h, speed(m/s)*3.6, or auto.",
    )
    parser.add_argument(
        "--quality-label",
        default=None,
        help="Optional curve quality label filter, for example 0 for normal curves only.",
    )
    parser.add_argument(
        "--cruise-min-speed-ratio",
        type=float,
        default=0.92,
        help="Cruise average uses speeds at least this ratio of the historical max speed.",
    )
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def format_float(value: float | None) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    if len(ordered) == 1:
        return ordered[0]
    pos = (len(ordered) - 1) * q
    lo = math.floor(pos)
    hi = math.ceil(pos)
    if lo == hi:
        return ordered[lo]
    return ordered[lo] * (hi - pos) + ordered[hi] * (pos - lo)


def mean(values: list[float]) -> float | None:
    if not values:
        return None
    return sum(values) / len(values)


def cruise_average(values: list[float], min_speed_ratio: float) -> tuple[float | None, int]:
    if not values:
        return None, 0
    max_speed = max(values)
    threshold = max_speed * min_speed_ratio
    cruise_values = [value for value in values if value >= threshold]
    return mean(cruise_values), len(cruise_values)


def find_col(headers: list[str], candidates: list[str]) -> int | None:
    normalized = {str(value).strip(): idx for idx, value in enumerate(headers)}
    for candidate in candidates:
        if candidate in normalized:
            return normalized[candidate]
    return None


def choose_speed(
    row: tuple[Any, ...],
    reference_idx: int | None,
    mps_idx: int | None,
    speed_source: str,
) -> float | None:
    ref_speed = as_float(row[reference_idx]) if reference_idx is not None else None
    mps_raw = as_float(row[mps_idx]) if mps_idx is not None else None
    mps_speed = mps_raw * 3.6 if mps_raw is not None else None
    if speed_source == "reference":
        return ref_speed
    if speed_source == "mps":
        return mps_speed
    return mps_speed if mps_speed is not None and mps_speed > 0 else ref_speed


def route_sections(route_map: Path, args: argparse.Namespace) -> list[str]:
    rows = read_csv(route_map)
    sections = [(row.get("section") or "").strip() for row in rows]
    sections = [section for section in sections if section]
    if args.sections:
        wanted = {item.strip() for item in args.sections.split(",") if item.strip()}
        return [section for section in sections if section in wanted]
    start = max(args.section_start, 1) - 1
    end = None if args.section_count is None else start + args.section_count
    return sections[start:end]


def discover_section_files(args: argparse.Namespace) -> list[tuple[int, str, Path]]:
    data_dir = Path(args.data_dir)
    if args.route_map:
        sections = route_sections(Path(args.route_map), args)
        return [
            (index + 1, section, data_dir / f"results_{section}.xlsx")
            for index, section in enumerate(sections)
        ]

    files = sorted(data_dir.glob("results_*.xlsx"), key=lambda p: p.name)
    if args.sections:
        wanted = {item.strip() for item in args.sections.split(",") if item.strip()}
        files = [path for path in files if path.stem.replace("results_", "", 1) in wanted]
    else:
        start = max(args.section_start, 1) - 1
        end = None if args.section_count is None else start + args.section_count
        files = files[start:end]
    return [
        (index + 1, path.stem.replace("results_", "", 1), path)
        for index, path in enumerate(files)
    ]


def summarize_section_file(
    section_index: int,
    section_name: str,
    xlsx_path: Path,
    args: argparse.Namespace,
) -> list[dict[str, Any]]:
    if not xlsx_path.exists():
        return [
            {
                "section_index": section_index,
                "section": section_name,
                "warning": "missing_section_xlsx",
                "xlsx_file": str(xlsx_path),
            }
        ]

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]

        section_idx = find_col(headers, SECTION_COLS)
        run_id_idx = find_col(headers, RUN_ID_COLS)
        segment_idx = find_col(headers, SEGMENT_COLS)
        quality_idx = find_col(headers, QUALITY_COLS)
        reference_idx = find_col(headers, REFERENCE_SPEED_COLS)
        mps_idx = find_col(headers, MPS_SPEED_COLS)
        if segment_idx is None:
            raise ValueError(f"{xlsx_path.name}: no segment column found.")
        if reference_idx is None and mps_idx is None:
            raise ValueError(f"{xlsx_path.name}: no speed column found.")

        groups: dict[tuple[int, str], dict[str, Any]] = {}

        for row in rows:
            segment_value = as_float(row[segment_idx])
            if segment_value is None:
                continue
            segment = int(segment_value)
            run_id = "" if run_id_idx is None or row[run_id_idx] is None else str(row[run_id_idx]).strip()
            quality = "" if quality_idx is None or row[quality_idx] is None else str(row[quality_idx]).strip()
            if args.quality_label is not None and quality != str(args.quality_label):
                continue
            speed = choose_speed(row, reference_idx, mps_idx, args.speed_source)
            if speed is None or speed < 0:
                continue

            key = (segment, run_id)
            group = groups.setdefault(key, {"speeds": [], "quality_labels": set()})
            group["speeds"].append(speed)
            if quality:
                group["quality_labels"].add(quality)

        output_rows: list[dict[str, Any]] = []
        for (segment, run_id), group in sorted(groups.items(), key=lambda item: (item[0][0], item[0][1])):
            speeds = group["speeds"]
            cruise_avg, cruise_count = cruise_average(speeds, args.cruise_min_speed_ratio)
            output_rows.append(
                {
                    "section_index": section_index,
                    "section": section_name,
                    "segment": segment,
                    "trip_no": segment + 1,
                    "run_id": run_id,
                    "quality_labels": "|".join(sorted(group["quality_labels"])),
                    "row_count": len(speeds),
                    "speed_source": args.speed_source,
                    "avg_speed_kmh": format_float(mean(speeds)),
                    "cruise_avg_speed_kmh": format_float(cruise_avg),
                    "cruise_point_count": cruise_count,
                    "cruise_min_speed_ratio": format_float(args.cruise_min_speed_ratio),
                    "max_speed_kmh": format_float(max(speeds) if speeds else None),
                    "p99_speed_kmh": format_float(percentile(speeds, 0.99)),
                    "p95_speed_kmh": format_float(percentile(speeds, 0.95)),
                    "xlsx_file": str(xlsx_path),
                    "warning": "",
                }
            )
        return output_rows
    finally:
        wb.close()


def main() -> int:
    args = parse_args()
    section_files = discover_section_files(args)
    if not section_files:
        raise ValueError("No section xlsx files selected.")

    all_rows: list[dict[str, Any]] = []
    for idx, (section_index, section, xlsx_path) in enumerate(section_files, 1):
        print(f"[{idx}/{len(section_files)}] {section}")
        rows = summarize_section_file(section_index, section, xlsx_path, args)
        all_rows.extend(rows)

    output_csv = Path(args.output_csv)
    write_csv(output_csv, all_rows)
    warnings = sum(1 for row in all_rows if row.get("warning"))
    print(f"Data dir:   {Path(args.data_dir)}")
    if args.route_map:
        print(f"Route map:  {Path(args.route_map)}")
    print(f"Output CSV: {output_csv}")
    print(f"Sections:   {len(section_files)}")
    print(f"Rows:       {len(all_rows)}")
    print(f"Warnings:   {warnings}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
