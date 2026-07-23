from __future__ import annotations

import argparse
import math
import statistics
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from plot_opentrack_tsvp_vs_history import (
    CUM_ENERGY_COL,
    DIST_COL,
    QUALITY_COL,
    SEGMENT_COL,
    SPEED_COL,
    STEP_ENERGY_COL,
    TIME_COL,
    append_history_section,
    as_float,
    energy_error,
    header_index,
    plot_compare,
    read_history_line,
    read_tsvp,
    selected_sections,
    write_csv,
)


PROJECT_ROOT = Path(r"D:\energy_conservation")
DEFAULT_TSVP_TEMPLATE = r"D:\OutPut\OT_priority_history_trip{trip:03d}.tsvP"
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_all_curve_quality"
DEFAULT_ROUTE_MAP = PROJECT_ROOT / "output" / "opentrack_route_map_newdoc" / "priority_dp_first5_route_map.csv"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output" / "opentrack_tsvp_compare" / "batch_history_trip001_070"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Batch plot OpenTrack .tsvP files against Step2 historical curves "
            "and summarize energy errors."
        )
    )
    parser.add_argument(
        "--trips",
        default="1-70",
        help="Trip numbers to process, for example 1-70 or 1,3,8-10.",
    )
    parser.add_argument(
        "--tsvp-template",
        default=DEFAULT_TSVP_TEMPLATE,
        help="OpenTrack .tsvP filename template. Use {trip:03d} or {trip}.",
    )
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Step2 results_*.xlsx directory.")
    parser.add_argument("--route-map", default=str(DEFAULT_ROUTE_MAP), help="Route map CSV for section order.")
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Batch output directory.")
    parser.add_argument("--prefix", default="priority_history", help="Prefix for output files.")
    parser.add_argument("--section-start", type=int, default=1, help="First 1-based section index to include.")
    parser.add_argument(
        "--section-count",
        type=int,
        default=5,
        help="Number of sections to include. Current OpenTrack course is usually first 5 sections.",
    )
    parser.add_argument(
        "--sections",
        default=None,
        help="Optional comma-separated section names. Overrides --section-start/--section-count.",
    )
    parser.add_argument(
        "--history-dwell",
        type=float,
        default=30.0,
        help="Dwell seconds inserted between historical sections for v-t/E-t comparison.",
    )
    parser.add_argument(
        "--quality-label",
        default=None,
        help="Optional historical curve quality filter, for example 0 for normal curves only.",
    )
    parser.add_argument(
        "--min-distance-ratio",
        type=float,
        default=0.90,
        help="Trips below this OpenTrack/history distance ratio are flagged and excluded from valid averages.",
    )
    parser.add_argument("--dpi", type=int, default=300, help="Output image DPI.")
    parser.add_argument("--image-format", choices=["png", "jpg"], default="jpg", help="Output image format.")
    parser.add_argument("--no-plot", action="store_true", help="Only calculate summaries, do not write images.")
    parser.add_argument("--xlsx", action="store_true", help="Also write an Excel summary workbook.")
    return parser.parse_args()


def parse_trip_selection(text: str) -> list[int]:
    trips: list[int] = []
    seen: set[int] = set()
    for raw_part in text.split(","):
        part = raw_part.strip()
        if not part:
            continue
        if "-" in part:
            left, right = [item.strip() for item in part.split("-", 1)]
            start = int(left)
            end = int(right)
            step = 1 if end >= start else -1
            numbers = range(start, end + step, step)
        else:
            numbers = [int(part)]
        for number in numbers:
            if number <= 0:
                raise ValueError(f"Trip number must be positive: {number}")
            if number not in seen:
                seen.add(number)
                trips.append(number)
    if not trips:
        raise ValueError("No trips selected.")
    return trips


def format_tsvp_path(template: str, trip: int) -> Path:
    try:
        return Path(template.format(trip=trip))
    except Exception as exc:
        raise ValueError(f"Invalid --tsvp-template {template!r}: {exc}") from exc


def pct_error(value: float | None, baseline: float | None) -> float | None:
    if value is None or baseline is None or abs(baseline) < 1e-9:
        return None
    return value / baseline * 100.0


def final_value(curve: dict[str, list[float]], key: str) -> float | None:
    values = curve.get(key) or []
    return values[-1] if values else None


def finite_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def summarize_values(values: list[float]) -> dict[str, float | str]:
    if not values:
        return {
            "count": 0,
            "mean": "",
            "mean_abs": "",
            "median": "",
            "median_abs": "",
            "min": "",
            "max": "",
            "max_abs": "",
            "rmse": "",
        }
    abs_values = [abs(value) for value in values]
    return {
        "count": len(values),
        "mean": statistics.fmean(values),
        "mean_abs": statistics.fmean(abs_values),
        "median": statistics.median(values),
        "median_abs": statistics.median(abs_values),
        "min": min(values),
        "max": max(values),
        "max_abs": max(abs_values),
        "rmse": math.sqrt(statistics.fmean([value * value for value in values])),
    }


def collect_metric(rows: list[dict[str, Any]], column: str, valid_only: bool) -> list[float]:
    values: list[float] = []
    for row in rows:
        if row.get("status") != "ok":
            continue
        if valid_only and str(row.get("valid_for_stats", "")).lower() != "true":
            continue
        number = finite_number(row.get(column))
        if number is not None:
            values.append(number)
    return values


def build_stats_rows(rows: list[dict[str, Any]], requested_count: int) -> list[dict[str, Any]]:
    ok_rows = [row for row in rows if row.get("status") == "ok"]
    valid_rows = [row for row in ok_rows if str(row.get("valid_for_stats", "")).lower() == "true"]
    failed_rows = [row for row in rows if row.get("status") != "ok"]
    flagged_rows = [row for row in ok_rows if str(row.get("valid_for_stats", "")).lower() != "true"]

    stats: list[dict[str, Any]] = [
        {"metric": "requested_trips", "scope": "all", "value": requested_count},
        {"metric": "plotted_or_calculated_trips", "scope": "all", "value": len(ok_rows)},
        {"metric": "valid_for_average_trips", "scope": "valid", "value": len(valid_rows)},
        {"metric": "flagged_distance_mismatch_trips", "scope": "all", "value": len(flagged_rows)},
        {"metric": "failed_trips", "scope": "all", "value": len(failed_rows)},
    ]

    for column in [
        "energy_error_pct",
        "energy_error_kwh",
        "duration_error_pct",
        "duration_error_s",
        "distance_error_pct",
        "distance_error_km",
    ]:
        for scope, valid_only in [("all_ok", False), ("valid_only", True)]:
            summary = summarize_values(collect_metric(rows, column, valid_only=valid_only))
            for name, value in summary.items():
                stats.append({"metric": f"{column}_{name}", "scope": scope, "value": value})

    total_ot = sum(finite_number(row.get("opentrack_energy_kwh")) or 0.0 for row in valid_rows)
    total_history = sum(finite_number(row.get("history_energy_kwh")) or 0.0 for row in valid_rows)
    total_diff = total_ot - total_history
    total_pct = pct_error(total_diff, total_history)
    stats.extend(
        [
            {"metric": "total_opentrack_energy_kwh", "scope": "valid_only", "value": total_ot},
            {"metric": "total_history_energy_kwh", "scope": "valid_only", "value": total_history},
            {"metric": "total_energy_error_kwh", "scope": "valid_only", "value": total_diff},
            {"metric": "total_energy_error_pct", "scope": "valid_only", "value": total_pct if total_pct is not None else ""},
        ]
    )
    return stats


def write_xlsx(path: Path, trip_rows: list[dict[str, Any]], stats_rows: list[dict[str, Any]]) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for --xlsx output.") from exc

    def write_sheet(ws: Any, rows: list[dict[str, Any]]) -> None:
        headers: list[str] = []
        for row in rows:
            for key in row:
                if key not in headers:
                    headers.append(key)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column_idx, header in enumerate(headers, start=1):
            width = min(max(len(str(header)) + 2, 12), 34)
            for cell in ws.iter_cols(min_col=column_idx, max_col=column_idx, min_row=2, max_row=ws.max_row):
                for item in cell:
                    width = min(max(width, len(str(item.value)) + 2), 42)
            ws.column_dimensions[get_column_letter(column_idx)].width = width

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "TripErrors"
    write_sheet(ws1, trip_rows)
    ws2 = wb.create_sheet("Stats")
    write_sheet(ws2, stats_rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_single_args(args: argparse.Namespace, trip: int) -> SimpleNamespace:
    return SimpleNamespace(
        data_dir=args.data_dir,
        trip_no=trip,
        run_id=None,
        quality_label=args.quality_label,
        history_dwell=args.history_dwell,
    )


def convert_history_records(records: list[tuple[float, float, float, float | None, float | None]]) -> dict[str, list[float]]:
    records.sort(key=lambda item: item[0])
    if not records:
        return {"time_s": [], "speed_kmh": [], "distance_km": [], "energy_kwh": []}

    t0 = records[0][0]
    s0 = records[0][2]
    first_cum = records[0][4]

    time_s: list[float] = []
    speed_kmh: list[float] = []
    distance_km: list[float] = []
    energy_kwh: list[float] = []
    running_energy_j = 0.0

    for t, v, s, e_step, e_cum in records:
        time_s.append(t - t0)
        speed_kmh.append(v)
        distance_km.append(s - s0)
        if e_cum is not None and first_cum is not None:
            energy_kwh.append((e_cum - first_cum) / 3_600_000.0)
        else:
            running_energy_j += e_step or 0.0
            energy_kwh.append(running_energy_j / 3_600_000.0)

    return {
        "time_s": time_s,
        "speed_kmh": speed_kmh,
        "distance_km": distance_km,
        "energy_kwh": energy_kwh,
    }


def read_history_section_many(
    xlsx_path: Path,
    trips: set[int],
    args: argparse.Namespace,
) -> dict[int, dict[str, list[float]]]:
    import openpyxl

    if not xlsx_path.exists():
        raise FileNotFoundError(f"Missing Step2 xlsx: {xlsx_path}")

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]

        idx_time = header_index(headers, TIME_COL)
        idx_speed = header_index(headers, SPEED_COL)
        idx_dist = header_index(headers, DIST_COL)
        idx_segment = header_index(headers, SEGMENT_COL)
        idx_quality = header_index(headers, QUALITY_COL)
        idx_energy = header_index(headers, STEP_ENERGY_COL)
        idx_cum_energy = header_index(headers, CUM_ENERGY_COL)

        required = {
            "time": idx_time,
            "speed": idx_speed,
            "distance": idx_dist,
            "segment": idx_segment,
        }
        missing = [name for name, idx in required.items() if idx is None]
        if missing:
            raise ValueError(f"{xlsx_path.name}: missing columns {missing}")
        if idx_cum_energy is None and idx_energy is None:
            raise ValueError(f"{xlsx_path.name}: missing energy/cumulative_energy column")

        records_by_trip: dict[int, list[tuple[float, float, float, float | None, float | None]]] = {
            trip: [] for trip in trips
        }
        min_segment = min(trips) - 1
        max_segment = max(trips) - 1

        for row in rows:
            segment = as_float(row[idx_segment]) if idx_segment is not None else None
            if segment is None:
                continue
            current_segment = int(segment)
            if current_segment > max_segment:
                break
            if current_segment < min_segment:
                continue
            trip = current_segment + 1
            if trip not in trips:
                continue

            if args.quality_label is not None and idx_quality is not None:
                quality = "" if row[idx_quality] is None else str(row[idx_quality]).strip()
                if quality != str(args.quality_label):
                    continue

            t = as_float(row[idx_time])
            v_ms = as_float(row[idx_speed])
            s_m = as_float(row[idx_dist])
            e_step = as_float(row[idx_energy]) if idx_energy is not None else None
            e_cum = as_float(row[idx_cum_energy]) if idx_cum_energy is not None else None
            if t is None or v_ms is None or s_m is None:
                continue
            records_by_trip[trip].append((t, v_ms * 3.6, s_m / 1000.0, e_step, e_cum))

        return {
            trip: convert_history_records(records)
            for trip, records in records_by_trip.items()
            if records
        }
    finally:
        wb.close()


def build_history_cache(
    sections: list[str],
    trips: list[int],
    args: argparse.Namespace,
) -> tuple[dict[int, dict[str, list[float]]], dict[int, list[str]]]:
    trip_set = set(trips)
    history_by_trip = {
        trip: {"time_s": [], "distance_km": [], "speed_kmh": [], "energy_kwh": []}
        for trip in trips
    }
    missing_by_trip: dict[int, list[str]] = {trip: [] for trip in trips}

    for section_idx, section in enumerate(sections):
        xlsx_path = Path(args.data_dir) / f"results_{section}.xlsx"
        print(f"  Loading history xlsx [{section_idx + 1}/{len(sections)}]: {xlsx_path.name}", flush=True)
        section_data_by_trip = read_history_section_many(xlsx_path, trip_set, args)
        for trip in trips:
            section_data = section_data_by_trip.get(trip)
            if section_data is None:
                missing_by_trip[trip].append(section)
                continue
            append_history_section(history_by_trip[trip], section_data, args.history_dwell, section_idx == 0)

    return history_by_trip, missing_by_trip


def process_trip(
    trip: int,
    args: argparse.Namespace,
    sections: list[str],
    plots_dir: Path,
    history_cache: dict[int, dict[str, list[float]]] | None = None,
    missing_history: dict[int, list[str]] | None = None,
) -> dict[str, Any]:
    tsvp_path = format_tsvp_path(args.tsvp_template, trip)
    row: dict[str, Any] = {
        "trip_no": trip,
        "tsvp_file": str(tsvp_path),
        "sections": len(sections),
        "history_dwell_s": args.history_dwell,
    }

    if not tsvp_path.exists():
        row.update({"status": "failed", "error": "missing_tsvp"})
        return row

    try:
        ot_curve = read_tsvp(tsvp_path)
        if history_cache is None:
            single_args = make_single_args(args, trip)
            history_curve, _history_summary = read_history_line(sections, single_args)
        else:
            missing_sections = (missing_history or {}).get(trip, [])
            if missing_sections:
                raise ValueError(f"missing historical section rows: {', '.join(missing_sections)}")
            history_curve = history_cache[trip]

        ot_duration = final_value(ot_curve, "time_s")
        history_duration = final_value(history_curve, "time_s")
        ot_distance = final_value(ot_curve, "distance_km")
        history_distance = final_value(history_curve, "distance_km")
        ot_energy = final_value(ot_curve, "energy_kwh")
        history_energy = final_value(history_curve, "energy_kwh")
        diff_kwh, diff_pct = energy_error(ot_curve, history_curve)

        duration_diff = (ot_duration - history_duration) if ot_duration is not None and history_duration is not None else None
        distance_diff = (ot_distance - history_distance) if ot_distance is not None and history_distance is not None else None
        duration_pct = pct_error(duration_diff, history_duration)
        distance_pct = pct_error(distance_diff, history_distance)
        distance_ratio = (ot_distance / history_distance) if ot_distance is not None and history_distance else None

        valid_for_stats = bool(
            diff_pct is not None
            and distance_ratio is not None
            and distance_ratio >= args.min_distance_ratio
        )
        flag = "ok" if valid_for_stats else "distance_mismatch_or_incomplete"

        output_plot = ""
        if not args.no_plot:
            output_plot_path = plots_dir / (
                f"{args.prefix}_trip{trip:03d}_vs_history_trip{trip:03d}_vt_vs_et_es.{args.image_format}"
            )
            title = (
                f"{args.prefix}_trip{trip:03d} vs history trip{trip:03d} | "
                f"{len(sections)} section(s) | history dwell={args.history_dwell:g}s"
            )
            plot_compare(ot_curve, history_curve, output_plot_path, title, args.dpi, show=False)
            output_plot = str(output_plot_path)

        row.update(
            {
                "status": "ok",
                "valid_for_stats": valid_for_stats,
                "flag": flag,
                "opentrack_points": len(ot_curve["time_s"]),
                "history_points": len(history_curve["time_s"]),
                "opentrack_duration_s": ot_duration,
                "history_duration_s": history_duration,
                "duration_error_s": duration_diff,
                "duration_error_pct": duration_pct,
                "opentrack_distance_km": ot_distance,
                "history_distance_km": history_distance,
                "distance_ratio": distance_ratio,
                "distance_error_km": distance_diff,
                "distance_error_pct": distance_pct,
                "opentrack_energy_kwh": ot_energy,
                "history_energy_kwh": history_energy,
                "energy_error_kwh": diff_kwh,
                "energy_error_pct": diff_pct,
                "plot": output_plot,
                "error": "",
            }
        )
    except Exception as exc:  # Keep the batch moving and report the bad trip in the table.
        row.update({"status": "failed", "valid_for_stats": False, "flag": "failed", "error": str(exc)})
    return row


def main() -> int:
    args = parse_args()
    trips = parse_trip_selection(args.trips)
    output_dir = Path(args.output_dir)
    plots_dir = output_dir / "plots"
    output_dir.mkdir(parents=True, exist_ok=True)
    if not args.no_plot:
        plots_dir.mkdir(parents=True, exist_ok=True)

    section_args = SimpleNamespace(
        section_start=args.section_start,
        section_count=args.section_count,
        sections=args.sections,
    )
    sections = selected_sections(Path(args.route_map), section_args)
    if not sections:
        raise ValueError("No sections selected from route map.")

    print(f"Trips:       {trips[0]}..{trips[-1]} ({len(trips)} selected)", flush=True)
    print(f"Sections:    {len(sections)} -> {', '.join(sections)}", flush=True)
    print(f"History dwell: {args.history_dwell:g}s", flush=True)
    print(f"Output dir:  {output_dir}", flush=True)
    print("Preloading historical curves...", flush=True)
    history_cache, missing_history = build_history_cache(sections, trips, args)

    trip_rows: list[dict[str, Any]] = []
    for index, trip in enumerate(trips, start=1):
        print(f"[{index}/{len(trips)}] trip{trip:03d}", flush=True)
        trip_rows.append(process_trip(trip, args, sections, plots_dir, history_cache, missing_history))

    stats_rows = build_stats_rows(trip_rows, requested_count=len(trips))
    detail_csv = output_dir / f"{args.prefix}_batch_trip_errors.csv"
    stats_csv = output_dir / f"{args.prefix}_batch_error_stats.csv"
    write_csv(detail_csv, trip_rows)
    write_csv(stats_csv, stats_rows)

    xlsx_path = output_dir / f"{args.prefix}_batch_error_summary.xlsx"
    if args.xlsx:
        write_xlsx(xlsx_path, trip_rows, stats_rows)

    ok_count = sum(1 for row in trip_rows if row.get("status") == "ok")
    valid_count = sum(1 for row in trip_rows if str(row.get("valid_for_stats", "")).lower() == "true")
    failed_count = len(trip_rows) - ok_count
    valid_energy_pcts = collect_metric(trip_rows, "energy_error_pct", valid_only=True)
    mean_abs = statistics.fmean([abs(value) for value in valid_energy_pcts]) if valid_energy_pcts else None
    mean_signed = statistics.fmean(valid_energy_pcts) if valid_energy_pcts else None

    print("", flush=True)
    print("Done.", flush=True)
    print(f"OK trips:     {ok_count}", flush=True)
    print(f"Valid trips:  {valid_count}", flush=True)
    print(f"Failed trips: {failed_count}", flush=True)
    if mean_abs is not None and mean_signed is not None:
        print(f"Mean energy error: {mean_signed:+.2f}%", flush=True)
        print(f"Mean abs error:    {mean_abs:.2f}%", flush=True)
    print(f"Detail CSV:   {detail_csv}", flush=True)
    print(f"Stats CSV:    {stats_csv}", flush=True)
    if args.xlsx:
        print(f"Excel:        {xlsx_path}", flush=True)
    if not args.no_plot:
        print(f"Plots:        {plots_dir}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
