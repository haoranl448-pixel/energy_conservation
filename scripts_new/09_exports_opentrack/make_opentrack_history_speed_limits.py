from __future__ import annotations

import argparse
import csv
import math
import time
from pathlib import Path
from typing import Any

import openpyxl

from opentrack_otd_client import print_response, send_command
from opentrack_otd_common import make_command_xml


PROJECT_ROOT = Path(r"D:\energy_conservation")
DEFAULT_DATA_DIR = PROJECT_ROOT / "data" / "data_processed_step2_v3_all_curve_quality"
DEFAULT_ROUTE_MAP = PROJECT_ROOT / "output" / "opentrack_route_map" / "priority_history_probe_route_map.csv"
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "opentrack_speed_limits" / "speed_limits_history.csv"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Extract historical per-section peak speeds from Step2 processed xlsx files "
            "and build OpenTrack setPositionSpeed commands."
        )
    )
    parser.add_argument("--data-dir", default=str(DEFAULT_DATA_DIR), help="Step2 processed xlsx directory.")
    parser.add_argument("--route-map", default=str(DEFAULT_ROUTE_MAP), help="OpenTrack route map CSV.")
    parser.add_argument("--output-csv", default=str(DEFAULT_OUTPUT), help="Historical speed limit table to write.")
    parser.add_argument(
        "--section-start",
        type=int,
        default=1,
        help="First 1-based section index to include.",
    )
    parser.add_argument(
        "--section-count",
        type=int,
        default=None,
        help="Only include this many sections from --section-start.",
    )
    parser.add_argument(
        "--sections",
        default=None,
        help="Optional comma-separated section names to include, overriding --section-start/--section-count.",
    )
    parser.add_argument(
        "--trip-no",
        type=int,
        default=1,
        help="Historical trip number, 1-based. Uses the xlsx segment column as trip_no - 1.",
    )
    parser.add_argument(
        "--run-id",
        default=None,
        help="Optional exact 日期+服务号. If provided, this overrides --trip-no.",
    )
    parser.add_argument(
        "--quality-label",
        default=None,
        help="Optional 曲线质量标签 filter, e.g. 0 for normal curves only.",
    )
    parser.add_argument(
        "--no-fast-segment-break",
        action="store_true",
        help=(
            "Disable early stopping after the requested segment block ends. "
            "Use this only if rows are not grouped by segment."
        ),
    )
    parser.add_argument(
        "--speed-source",
        choices=["reference", "mps", "auto"],
        default="auto",
        help=(
            "Speed column to use. reference=BCU6_CCU 参考速度km/h, "
            "mps=速度(m/s)*3.6, auto prefers mps when available."
        ),
    )
    parser.add_argument(
        "--method",
        choices=["max", "p99", "p95"],
        default="max",
        help="How to summarize the historical speed curve into a section speed limit.",
    )
    parser.add_argument(
        "--speed-margin",
        type=float,
        default=0.0,
        help="Additional km/h added after extracting the historical speed.",
    )
    parser.add_argument("--ceil-speed", action="store_true", help="Round the final speed up to whole km/h.")
    parser.add_argument(
        "--range-mode",
        choices=["single-route", "next-route-boundary"],
        default="single-route",
        help="OpenTrack route range style. single-route is recommended with routeEntry fallback maps.",
    )
    parser.add_argument(
        "--single-route-end-offset",
        type=float,
        default=999999.0,
        help="End offset used by --range-mode single-route.",
    )
    parser.add_argument(
        "--missing-end-offset",
        type=float,
        default=999999.0,
        help="Fallback end offset when route map lacks a precise endRouteOffset.",
    )
    parser.add_argument(
        "--train-id",
        default=None,
        help="Optional OpenTrack trainID. For historical comparison this is often priority_history.",
    )
    parser.add_argument("--send", action="store_true", help="Send setPositionSpeed commands to OpenTrack.")
    parser.add_argument("--host", default="127.0.0.1", help="OpenTrack OTD server host.")
    parser.add_argument("--port", type=int, default=9002, help="OpenTrack OTD server port.")
    parser.add_argument("--path", default="/otd", help="OTD HTTP path.")
    parser.add_argument("--timeout", type=float, default=10.0, help="Socket timeout in seconds.")
    parser.add_argument("--wait-response", action="store_true", help="Wait for HTTP responses from OpenTrack.")
    parser.add_argument("--sleep", type=float, default=0.02, help="Delay between sent commands.")
    parser.add_argument("--verbose", action="store_true", help="Print SOAP command payloads while sending.")
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames: list[str] = []
    for row in rows:
        for key in row:
            if key not in fieldnames:
                fieldnames.append(key)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def as_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        value_f = float(text)
    except ValueError:
        return None
    if math.isnan(value_f) or math.isinf(value_f):
        return None
    return value_f


def format_float(value: float | None) -> str:
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    values = sorted(values)
    if len(values) == 1:
        return values[0]
    pos = (len(values) - 1) * q
    lo = math.floor(pos)
    hi = math.ceil(pos)
    if lo == hi:
        return values[lo]
    return values[lo] * (hi - pos) + values[hi] * (pos - lo)


def summarize_speed(values: list[float], method: str) -> float | None:
    if not values:
        return None
    if method == "p99":
        return percentile(values, 0.99)
    if method == "p95":
        return percentile(values, 0.95)
    return max(values)


def find_col(headers: list[str], candidates: list[str]) -> int | None:
    normalized = {str(name).strip(): idx for idx, name in enumerate(headers)}
    for candidate in candidates:
        if candidate in normalized:
            return normalized[candidate]
    return None


def choose_speed(
    row: tuple[Any, ...],
    reference_idx: int | None,
    mps_idx: int | None,
    speed_source: str,
) -> float | None:
    ref_speed = as_float(row[reference_idx]) if reference_idx is not None else None
    mps_speed = as_float(row[mps_idx]) * 3.6 if mps_idx is not None and as_float(row[mps_idx]) is not None else None
    if speed_source == "reference":
        return ref_speed
    if speed_source == "mps":
        return mps_speed
    return mps_speed if mps_speed is not None and mps_speed > 0 else ref_speed


def extract_section_history_speed(
    xlsx_path: Path,
    trip_no: int,
    run_id: str | None,
    quality_label: str | None,
    speed_source: str,
    method: str,
    fast_segment_break: bool,
) -> dict[str, Any]:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]

    section_idx = find_col(headers, ["区段", "站间区间", "section"])
    run_id_idx = find_col(headers, ["日期+服务号"])
    segment_idx = find_col(headers, ["segment"])
    quality_idx = find_col(headers, ["曲线质量标签"])
    reference_idx = find_col(headers, ["BCU6_CCU 参考速度km/h", "参考速度km/h"])
    mps_idx = find_col(headers, ["速度(m/s)"])

    if reference_idx is None and mps_idx is None:
        raise ValueError(f"{xlsx_path.name}: no speed column found.")
    if run_id is None and segment_idx is None:
        raise ValueError(f"{xlsx_path.name}: no segment column found; use --run-id.")

    target_segment = trip_no - 1
    speeds: list[float] = []
    selected_run_ids: set[str] = set()
    selected_segments: set[str] = set()
    selected_quality: set[str] = set()
    section_name = ""
    seen_target_segment = False

    for row in rows:
        if section_idx is not None and not section_name and row[section_idx] is not None:
            section_name = str(row[section_idx]).strip()

        if run_id is not None:
            current_run_id = "" if run_id_idx is None or row[run_id_idx] is None else str(row[run_id_idx]).strip()
            if current_run_id != run_id:
                continue
        else:
            current_segment = as_float(row[segment_idx]) if segment_idx is not None else None
            if (
                fast_segment_break
                and seen_target_segment
                and current_segment is not None
                and int(current_segment) != target_segment
            ):
                break
            if current_segment is None or int(current_segment) != target_segment:
                continue
            seen_target_segment = True

        if quality_label is not None and quality_idx is not None:
            current_quality = "" if row[quality_idx] is None else str(row[quality_idx]).strip()
            if current_quality != str(quality_label):
                continue

        speed = choose_speed(row, reference_idx, mps_idx, speed_source)
        if speed is None or speed < 0:
            continue
        speeds.append(speed)
        if run_id_idx is not None and row[run_id_idx] is not None:
            selected_run_ids.add(str(row[run_id_idx]).strip())
        if segment_idx is not None and row[segment_idx] is not None:
            selected_segments.add(str(row[segment_idx]).strip())
        if quality_idx is not None and row[quality_idx] is not None:
            selected_quality.add(str(row[quality_idx]).strip())

    speed_limit = summarize_speed(speeds, method)
    return {
        "section": section_name or xlsx_path.stem.replace("results_", "", 1),
        "xlsx_file": str(xlsx_path),
        "row_count": len(speeds),
        "run_ids": "|".join(sorted(selected_run_ids)),
        "segments": "|".join(sorted(selected_segments)),
        "quality_labels": "|".join(sorted(selected_quality)),
        "historical_speed_kmh": speed_limit,
        "max_speed_kmh": max(speeds) if speeds else None,
        "p99_speed_kmh": percentile(speeds, 0.99),
        "p95_speed_kmh": percentile(speeds, 0.95),
    }


def build_route_index(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    index: dict[str, dict[str, str]] = {}
    for row in rows:
        section = (row.get("section") or row.get("站间区间") or "").strip()
        if section:
            index[section] = row
    return index


def make_limit_rows(args: argparse.Namespace) -> list[dict[str, Any]]:
    route_rows = read_csv(Path(args.route_map))
    route_index = build_route_index(route_rows)
    data_dir = Path(args.data_dir)
    output_rows: list[dict[str, Any]] = []
    selected_sections = None
    if args.sections:
        selected_sections = {item.strip() for item in args.sections.split(",") if item.strip()}

    for order, route in enumerate(route_rows, 1):
        section = (route.get("section") or "").strip()
        if not section:
            continue
        if selected_sections is not None:
            if section not in selected_sections:
                continue
        elif order < args.section_start:
            continue
        elif args.section_count is not None and order >= args.section_start + args.section_count:
            continue
        xlsx_path = data_dir / f"results_{section}.xlsx"
        warnings: list[str] = []
        if not xlsx_path.exists():
            warnings.append("missing_section_xlsx")
            hist = {
                "section": section,
                "xlsx_file": str(xlsx_path),
                "row_count": 0,
                "run_ids": "",
                "segments": "",
                "quality_labels": "",
                "historical_speed_kmh": None,
                "max_speed_kmh": None,
                "p99_speed_kmh": None,
                "p95_speed_kmh": None,
            }
        else:
            hist = extract_section_history_speed(
                xlsx_path=xlsx_path,
                trip_no=args.trip_no,
                run_id=args.run_id,
                quality_label=args.quality_label,
                speed_source=args.speed_source,
                method=args.method,
                fast_segment_break=not args.no_fast_segment_break,
            )
            if not hist["row_count"]:
                warnings.append("no_matching_rows")

        speed_kmh = hist["historical_speed_kmh"]
        if speed_kmh is None:
            warnings.append("missing_historical_speed")
            speed_limit = None
        else:
            speed_limit = float(speed_kmh) + args.speed_margin
            if args.ceil_speed:
                speed_limit = float(math.ceil(speed_limit))

        current_route = route_index.get(section, route)
        start_route_id = (current_route.get("startRouteID") or "").strip()
        start_offset = as_float(current_route.get("startRouteOffset"))
        end_route_id = (current_route.get("endRouteID") or start_route_id).strip()
        end_offset = as_float(current_route.get("endRouteOffset"))

        if start_route_id and start_offset is None:
            start_offset = 0.0
            warnings.append("start_offset_defaulted_to_0")
        if end_route_id and end_offset is None:
            end_offset = args.missing_end_offset
            warnings.append("end_offset_defaulted")
        if args.range_mode == "single-route" and start_route_id:
            end_route_id = start_route_id
            end_offset = args.single_route_end_offset

        attrs: dict[str, Any] = {}
        if args.train_id:
            attrs["trainID"] = args.train_id
        if speed_limit is not None:
            attrs["speed"] = speed_limit
        attrs["startRouteID"] = start_route_id
        attrs["startRouteOffset"] = start_offset
        attrs["endRouteID"] = end_route_id
        attrs["endRouteOffset"] = end_offset
        xml = make_command_xml("setPositionSpeed", attrs) if speed_limit is not None and start_route_id else ""

        output_rows.append(
            {
                "send_order": order,
                "section": section,
                "trip_no": args.trip_no,
                "run_id_filter": args.run_id or "",
                "matched_run_ids": hist["run_ids"],
                "matched_segments": hist["segments"],
                "quality_labels": hist["quality_labels"],
                "row_count": hist["row_count"],
                "method": args.method,
                "historical_speed_kmh": format_float(speed_kmh),
                "max_speed_kmh": format_float(hist["max_speed_kmh"]),
                "p99_speed_kmh": format_float(hist["p99_speed_kmh"]),
                "p95_speed_kmh": format_float(hist["p95_speed_kmh"]),
                "speed_margin_kmh": format_float(args.speed_margin),
                "speed_limit_kmh": format_float(speed_limit),
                "range_mode": args.range_mode,
                "startRouteID": start_route_id,
                "startRouteOffset": format_float(start_offset),
                "endRouteID": end_route_id,
                "endRouteOffset": format_float(end_offset),
                "trainID": args.train_id or "",
                "otd_xml": xml,
                "warning": ";".join(warnings),
            }
        )
    return output_rows


def send_rows(rows: list[dict[str, Any]], args: argparse.Namespace) -> int:
    failures = 0
    sent = 0
    for row in rows:
        speed = as_float(row.get("speed_limit_kmh"))
        if speed is None or not row.get("startRouteID"):
            failures += 1
            print(f"skip {row.get('section')}: {row.get('warning')}")
            continue
        attrs: dict[str, Any] = {}
        if args.train_id:
            attrs["trainID"] = args.train_id
        attrs["speed"] = speed
        attrs["startRouteID"] = row.get("startRouteID")
        attrs["startRouteOffset"] = as_float(row.get("startRouteOffset"))
        attrs["endRouteID"] = row.get("endRouteID")
        attrs["endRouteOffset"] = as_float(row.get("endRouteOffset"))
        status, text = send_command("setPositionSpeed", attrs, args)
        sent += 1
        print(f"[{sent:02d}] {row.get('section')} history <= {row.get('speed_limit_kmh')} km/h")
        print_response(status, text)
        if not (status == 0 or 200 <= status < 300):
            failures += 1
        if args.sleep > 0:
            time.sleep(args.sleep)
    print(f"Sent {sent} historical setPositionSpeed command(s), failures={failures}")
    return 0 if failures == 0 else 1


def main() -> int:
    args = parse_args()
    rows = make_limit_rows(args)
    write_csv(Path(args.output_csv), rows)

    warnings = sum(1 for row in rows if row.get("warning"))
    print(f"Data dir:   {Path(args.data_dir)}")
    print(f"Route map:  {Path(args.route_map)}")
    print(f"Output CSV: {Path(args.output_csv)}")
    print(f"Trip no:    {args.trip_no}")
    if args.run_id:
        print(f"Run ID:     {args.run_id}")
    print(f"Sections:   {len(rows)}")
    if args.sections:
        print(f"Filter:     sections={args.sections}")
    elif args.section_count is not None or args.section_start != 1:
        print(f"Filter:     section_start={args.section_start}, section_count={args.section_count}")
    print(f"Warnings:   {warnings}")
    if args.send:
        return send_rows(rows, args)
    print("Not sent. Add --send to push historical setPositionSpeed commands to OpenTrack.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
